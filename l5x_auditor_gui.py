#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
L5X Auditor - análisis estático de proyectos Studio 5000 / Logix Designer (.L5X)

Autor: Joetan Saldaña
Copyright (c) 2026 Joetan Saldaña. Todos los derechos reservados.

Busca patrones que normalmente indican lógica puenteada, deshabilitada,
forzada o mal escrita. NO sustituye una revisión manual: cada hallazgo
hay que confirmarlo en el contexto de la máquina.

Uso:
    python l5x_auditor_gui.py                          # interfaz grafica
    python l5x_auditor_gui.py proyecto.L5X             # resumen en consola
    python l5x_auditor_gui.py *.L5X --csv r.csv --html r.html --json r.json
    python l5x_auditor_gui.py proyecto.L5X --min-sev MEDIA --fail-on ALTA
    python l5x_auditor_gui.py --list-rules
    python l5x_auditor_gui.py --gui proyecto.L5X       # abre la GUI con archivos

Solo usa la biblioteca estándar de Python. Opcionales:
    openpyxl     -> exportar a Excel (.xlsx)
    tkinterdnd2  -> arrastrar y soltar archivos sobre la ventana
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import os
import queue
import re
import sys
import tempfile
import threading
import time
import webbrowser
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime

__version__ = "2.0.0"
__author__ = "Joetan Saldaña"
APP_NAME = "L5X Auditor"
APP_TITLE = "L5X Auditor - Auditoría de proyectos Studio 5000"
APP_COPYRIGHT = f"© 2026 {__author__}"
APP_CREDIT = f"Desarrollado por {__author__}"


def resource_path(*parts):
    """Ruta a un recurso, tanto en el script como dentro del .exe de PyInstaller."""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, *parts)

SEVERITIES = ("ALTA", "MEDIA", "BAJA", "INFO")
SEVERITY_ORDER = {"ALTA": 0, "MEDIA": 1, "BAJA": 2, "INFO": 3}
SEV_WEIGHT = {"ALTA": 10, "MEDIA": 3, "BAJA": 1, "INFO": 0}
STATUSES = ("Pendiente", "Confirmado", "Falso positivo", "Aceptado")


# ----------------------------------------------------------------- reglas -----
@dataclass(frozen=True)
class Rule:
    code: str
    sev: str
    title: str
    category: str
    desc: str
    fix: str


RULES = {r.code: r for r in (
    Rule("RAMA-SIEMPRE-VERDADERA", "ALTA", "Rama siempre verdadera", "Lógica",
         "XIC y XIO del mismo tag en ramas paralelas: una de las dos siempre conduce, "
         "así que la rama no condiciona nada. Es la forma típica de puentear un interlock.",
         "Identificar qué condición debía evaluar la rama y restaurar el interlock original."),
    Rule("SERIE-SIEMPRE-FALSA", "ALTA", "Serie siempre falsa", "Lógica",
         "XIC y XIO del mismo tag en el mismo camino serie: el renglón (o la rama) "
         "nunca puede conducir.",
         "Confirmar si es una forma deliberada de deshabilitar lógica; si lo es, "
         "documentarlo o eliminar el renglón."),
    Rule("VALOR-FIJO", "ALTA", "Constante escrita en un resultado", "Datos",
         "Se escribe una constante en un tag de resultado de inspección o medición. "
         "Puede ocultar un rechazo real y dejar pasar producto malo.",
         "Verificar que la escritura solo ocurra en modo prueba / dry cycle y que no "
         "pueda activarse en producción."),
    Rule("ONS-DUPLICADO", "ALTA", "Bit de one-shot reutilizado", "Lógica",
         "El mismo bit de almacenamiento se usa en varios ONS/OSR/OSF. Los flancos "
         "se interfieren entre si y se pierden o repiten pulsos.",
         "Asignar un bit de almacenamiento único a cada instrucción de flanco."),
    Rule("ESCRITURA-ENTRADA", "ALTA", "Escritura sobre entrada física", "E/S",
         "Una instrucción escribe sobre la imagen de entradas (:I.). El valor real del "
         "campo queda enmascarado o se sobrescribe en la siguiente actualización.",
         "Usar un tag interno para simular la señal; nunca escribir sobre la imagen de entradas."),
    Rule("JSR-INEXISTENTE", "ALTA", "JSR a rutina inexistente", "Estructura",
         "Se llama con JSR/FOR a una rutina que no existe en el programa.",
         "Corregir el nombre de la rutina o crearla; el proyecto no verificará."),
    Rule("FORZADO", "ALTA", "Datos de forzado en el export", "E/S",
         "El tag contiene una máscara de forzado distinta de cero. Un force activo "
         "sustituye el valor real de la E/S.",
         "Revisar y retirar los forces antes de liberar el proyecto."),
    Rule("AFI", "MEDIA", "Lógica deshabilitada con AFI", "Lógica",
         "AFI() fuerza el renglón a falso. Es común durante pruebas pero no debería "
         "quedar en producción sin documentar.",
         "Eliminar el AFI o documentar en el comentario por qué la lógica está deshabilitada."),
    Rule("TIMER-PRE-0", "MEDIA", "Temporizador con preset 0", "Lógica",
         "El preset del temporizador se carga con 0: el temporizador termina de inmediato "
         "y el retardo queda anulado.",
         "Confirmar el valor de preset esperado y quitar el MOV(0, ...PRE)."),
    Rule("OTE-DUPLICADA", "MEDIA", "Bobina duplicada", "Lógica",
         "La misma OTE se escribe en varios renglones; solo cuenta el último que se "
         "ejecuta (doble bobina).",
         "Unificar la lógica en un solo renglón con ramas en paralelo."),
    Rule("OTL-SIN-OTU", "MEDIA", "OTL sin OTU", "Lógica",
         "El tag se enclava con OTL pero no se ve un OTU en el mismo programa. "
         "Puede quedar enclavado para siempre (o se libera desde otro programa / HMI).",
         "Verificar dónde se libera el enclavamiento. Suele ser normal en fallas que se "
         "resetean desde HMI."),
    Rule("TIMER-DUPLICADO", "MEDIA", "Temporizador reutilizado", "Lógica",
         "El mismo tag de temporizador se usa en varias instrucciones TON/TOF/RTO. "
         "Los acumulados se pisan entre si.",
         "Usar un temporizador independiente para cada instrucción."),
    Rule("TAREA-INHIBIDA", "MEDIA", "Tarea inhibida", "Estructura",
         "La tarea está marcada como inhibida: ninguno de sus programas se ejecuta.",
         "Confirmar que la inhibición es intencional."),
    Rule("PROGRAMA-DESHABILITADO", "MEDIA", "Programa deshabilitado", "Estructura",
         "El programa está deshabilitado y no se ejecuta.",
         "Confirmar que es intencional o eliminar el programa."),
    Rule("ST-CONDICION-CONSTANTE", "MEDIA", "Condición constante en ST", "Lógica",
         "Un IF/ELSIF/WHILE de texto estructurado evalúa una constante (TRUE, FALSE, 1, 0): "
         "el bloque siempre o nunca se ejecuta.",
         "Restaurar la condición real o eliminar el bloque muerto."),
    Rule("JMP", "BAJA", "Salto JMP", "Estructura",
         "Uso de JMP: los renglones entre el JMP y su LBL se saltan cuando la condición es verdadera.",
         "Revisar que el salto no omita lógica de seguridad o de interlock."),
    Rule("COMENTARIO-SOSPECHOSO", "BAJA", "Comentario sospechoso", "Documentación",
         "El comentario menciona puentes, bypass, pruebas o cambios temporales.",
         "Leer el comentario y confirmar si la modificación temporal sigue activa."),
    Rule("RUTINA-SIN-LLAMAR", "BAJA", "Rutina nunca llamada", "Estructura",
         "Ninguna instrucción JSR/FOR del programa llama a esta rutina y no es la rutina "
         "principal ni la de fallas: su lógica nunca se ejecuta.",
         "Eliminar la rutina o agregar el JSR faltante."),
    Rule("PROGRAMA-SIN-TAREA", "BAJA", "Programa sin programar", "Estructura",
         "El programa no está asignado a ninguna tarea, así que nunca se ejecuta.",
         "Asignarlo a una tarea o eliminarlo."),
    Rule("DEBUG-BIT", "INFO", "Bit de debug / bypass", "Datos",
         "Un bit de debug o bypass condiciona lógica de producción. Se muestra su valor "
         "actual en el export.",
         "Confirmar que el bit está en el valor seguro para producción."),
)}


# ---------------------------------------------------------------- parser -----
_INST_RE = re.compile(r"([A-Za-z_][A-Za-z0-9_]*)\s*\(")


def split_top(text, sep=","):
    """Separa por 'sep' respetando parentesis y corchetes."""
    parts, depth, cur = [], 0, []
    for ch in text:
        if ch in "([":
            depth += 1
        elif ch in ")]":
            depth -= 1
        if ch == sep and depth == 0:
            parts.append("".join(cur))
            cur = []
        else:
            cur.append(ch)
    parts.append("".join(cur))
    return parts


def parse_series(text):
    """Convierte el texto de un rung en una lista.
    Instrucción -> ('inst', nombre, [args])
    Rama        -> ('branch', [serie, serie, ...])
    """
    items, i, n = [], 0, len(text)
    while i < n:
        ch = text[i]
        if ch in " \t\r\n;":
            i += 1
            continue
        if ch == "[":
            depth, j = 1, i + 1
            while j < n and depth:
                if text[j] == "[":
                    depth += 1
                elif text[j] == "]":
                    depth -= 1
                j += 1
            inner = text[i + 1:j - 1]
            items.append(("branch", [parse_series(p) for p in split_top(inner)]))
            i = j
            continue
        m = _INST_RE.match(text, i)
        if m:
            name, start = m.group(1), m.end()
            j, depth = start, 1
            while j < n and depth:
                if text[j] == "(":
                    depth += 1
                elif text[j] == ")":
                    depth -= 1
                j += 1
            args = [a.strip() for a in split_top(text[start:j - 1])]
            items.append(("inst", name, args))
            i = j
            continue
        i += 1
    return items


def walk(series):
    """Recorre recursivamente todas las instrucciones de una serie."""
    for it in series:
        if it[0] == "inst":
            yield it
        else:
            for sub in it[1]:
                yield from walk(sub)


def contacts_of(series, kind):
    """Contactos 'kind' (XIC/XIO) directos en este nivel de serie."""
    return [it[2][0] for it in series if it[0] == "inst" and it[1] == kind and it[2]]


# --------------------------------------------------------------- lectura -----
@dataclass
class Project:
    path: str
    info: dict = field(default_factory=dict)
    programs: dict = field(default_factory=dict)
    tasks: list = field(default_factory=list)
    rungs: list = field(default_factory=list)
    st_lines: list = field(default_factory=list)
    values: dict = field(default_factory=dict)
    forces: list = field(default_factory=list)
    tag_count: int = 0

    @property
    def scheduled(self):
        return {p for t in self.tasks for p in t["programs"]}

    @property
    def full_export(self):
        return self.info.get("TargetType", "Controller") == "Controller"


def _has_force(text):
    body = re.sub(r"\b\d+#", "", text or "")
    return bool(re.search(r"[1-9A-Fa-f]", body))


def load_project(path):
    """Lee un L5X y devuelve un Project."""
    root = ET.parse(path).getroot()
    if root.tag != "RSLogix5000Content":
        raise ValueError(f"No parece un L5X de Studio 5000 (raiz <{root.tag}>)")
    proj = Project(path=path)
    for k in ("SoftwareRevision", "TargetName", "TargetType", "ExportDate"):
        if root.get(k):
            proj.info[k] = root.get(k)

    def scan_tags(scope_el, scope_name):
        for tag in scope_el.iter("Tag"):
            name = tag.get("Name")
            proj.tag_count += 1
            for data in tag:
                if data.tag == "Data" and data.get("Format") == "L5K" and data.text:
                    proj.values[(scope_name, name)] = data.text.strip()
                elif "Force" in data.tag and _has_force(data.text):
                    proj.forces.append((scope_name, name))

    ctrl = root.find(".//Controller")
    if ctrl is not None:
        for k in ("Name", "ProcessorType", "MajorRev", "MinorRev", "LastModifiedDate"):
            if ctrl.get(k):
                proj.info[k] = ctrl.get(k)
        tags_el = ctrl.find("Tags")
        if tags_el is not None:
            scan_tags(tags_el, "<controller>")

    for prog in root.iter("Program"):
        pname = prog.get("Name")
        p = proj.programs.setdefault(pname, {
            "main": prog.get("MainRoutineName"), "fault": prog.get("FaultRoutineName"),
            "disabled": prog.get("Disabled", "false").lower() == "true",
            "parent": prog.get("Parent"), "class": prog.get("Class", "Standard"),
            "routines": {}, "rungs": 0, "st": 0,
        })
        tags_el = prog.find("Tags")
        if tags_el is not None:
            scan_tags(tags_el, pname)
        for rout in prog.iter("Routine"):
            rname = rout.get("Name")
            p["routines"][rname] = rout.get("Type", "RLL")
            for rung in rout.iter("Rung"):
                t = rung.find("Text")
                if t is None or not t.text:
                    continue
                p["rungs"] += 1
                proj.rungs.append({
                    "program": pname,
                    "routine": rname,
                    "number": rung.get("Number"),
                    "text": t.text.strip(),
                    "comment": (rung.findtext("Comment") or "").strip(),
                })
            st = rout.find("STContent")
            if st is not None:
                for line in st.iter("Line"):
                    if line.text and line.text.strip():
                        p["st"] += 1
                        proj.st_lines.append({
                            "program": pname, "routine": rname,
                            "number": line.get("Number"), "text": line.text.rstrip(),
                        })

    for task in root.iter("Task"):
        proj.tasks.append({
            "name": task.get("Name"), "type": task.get("Type", ""),
            "rate": task.get("Rate"), "priority": task.get("Priority"),
            "inhibit": task.get("InhibitTask", "false").lower() == "true",
            "programs": [sp.get("Name") for sp in task.iter("ScheduledProgram")],
        })
    return proj


def load(path):
    """Compatibilidad con la version anterior: devuelve (rungs, tag_values)."""
    proj = load_project(path)
    return proj.rungs, proj.values


# --------------------------------------------------------------- chequeos ----
_SUSP_RE = re.compile(
    r"\b(bypass\w*|by-pass|puente\w*|puentead\w*|jumper\w*|temporal\w*|provisional\w*|"
    r"prueba\w*|test\w*|forzad\w*|override|hack\w*|quitar|remover|deshabilitad\w*)\b", re.I)
_INPUT_RE = re.compile(r":I\.")
_DRY_RE = re.compile(r"dry_?cycle", re.I)
_DEBUG_RE = re.compile(r"^g?[Dd]ebug|_Debug|bypass", re.I)
# Destino segun instruccion (indice del argumento que se escribe)
_DEST_ARG = {"OTE": 0, "OTL": 0, "OTU": 0, "CLR": 0, "CPT": 0, "MOV": 1, "COP": 1,
             "CPS": 1, "FLL": 1, "ADD": 2, "SUB": 2, "MUL": 2, "DIV": 2}


def _finding(code, prog, rout, rung, msg, text="", sev=None, **extra):
    d = {"sev": sev or RULES[code].sev, "code": code, "prog": prog or "-",
         "rout": rout or "-", "rung": "-" if rung is None else str(rung),
         "msg": msg, "text": (text or "")[:4000]}
    d.update(extra)
    return d


def _valor_fijo_msg(src, dst, dry):
    if not re.fullmatch(r"-?\d+(\.\d+)?", src) or float(src) == 0:
        return None
    if re.search(r"\.(ID|PRE|ACC|Len|LEN)\b|\.Lim\.|Check\[", dst):
        return None
    if not re.search(r"(Status_Res|\.Res\[|Results\[|Inspect|Measure)", dst, re.I):
        return None
    return (f"Escribe la constante {src} en '{dst}'" +
            (" (rama de dry cycle)" if dry else ": posible resultado forzado"))


def check_rung(r, findings):
    text = r["text"]
    series = parse_series(text)
    P, R, N = r["program"], r["routine"], r["number"]

    def add(code, msg, sev=None, **kw):
        findings.append(_finding(code, P, R, N, msg, text, sev, **kw))

    # 1. Rama siempre verdadera: XIC(t) y XIO(t) en ramas hermanas.
    def scan_branches(s):
        for it in s:
            if it[0] == "branch":
                xic, xio = set(), set()
                for leg in it[1]:
                    if len(leg) == 1 and leg[0][0] == "inst":
                        nm, args = leg[0][1], leg[0][2]
                        if nm == "XIC" and args:
                            xic.add(args[0])
                        if nm == "XIO" and args:
                            xio.add(args[0])
                for t in sorted(xic & xio):
                    add("RAMA-SIEMPRE-VERDADERA",
                        f"XIC y XIO de '{t}' en ramas paralelas: el interlock queda anulado")
                for leg in it[1]:
                    scan_branches(leg)

    scan_branches(series)

    # 2. Serie siempre falsa: XIC(t) y XIO(t) en el mismo camino serie.
    def scan_series(s):
        xic, xio = set(contacts_of(s, "XIC")), set(contacts_of(s, "XIO"))
        for t in sorted(xic & xio):
            add("SERIE-SIEMPRE-FALSA",
                f"XIC y XIO de '{t}' en serie: el renglón nunca puede conducir")
        for it in s:
            if it[0] == "branch":
                for leg in it[1]:
                    scan_series(leg)

    scan_series(series)

    # 3. AFI: logica deshabilitada a proposito.
    if re.search(r"\bAFI\s*\(\s*\)", text):
        add("AFI", "Contiene AFI(): hay lógica deshabilitada en el renglón")

    dry = bool(_DRY_RE.search(text))
    for it in walk(series):
        name, args = it[1], it[2]
        if not args:
            continue
        # 4. Valores fijos escritos a resultados de inspeccion / medicion.
        if name in ("MOV", "FLL") and len(args) >= 2:
            msg = _valor_fijo_msg(args[0], args[1], dry)
            if msg:
                add("VALOR-FIJO", msg, "BAJA" if dry else None)
        # 5. Temporizadores con preset 0.
        if name in ("TON", "TOF", "RTO"):
            if re.search(r"MOV\(\s*0\s*,\s*" + re.escape(args[0]) + r"\.PRE\b", text):
                add("TIMER-PRE-0", f"Preset de '{args[0]}' se carga con 0")
        # 6. Bits de debug / bypass que condicionan logica.
        if name in ("XIC", "XIO") and _DEBUG_RE.search(args[0]):
            add("DEBUG-BIT", f"Condicionado por '{args[0]}'", tag=args[0])
        # 7. Saltos.
        if name == "JMP":
            add("JMP", f"JMP({args[0]}): revisar que no salte lógica de seguridad")
        # 8. Escritura sobre entradas fisicas.
        idx = _DEST_ARG.get(name)
        if idx is not None and idx < len(args) and _INPUT_RE.search(args[idx]):
            add("ESCRITURA-ENTRADA", f"{name} escribe sobre la entrada '{args[idx]}'")

    # 9. Comentario que habla de puentes / pruebas.
    m = _SUSP_RE.search(r.get("comment") or "")
    if m:
        c = " ".join(r["comment"].split())
        add("COMENTARIO-SOSPECHOSO", f"Menciona '{m.group(0)}': {c[:140]}")

    return series


_ST_ASSIGN_RE = re.compile(r"([A-Za-z_][\w\.\[\]:]*)\s*:=\s*([^;]+);")
_ST_CONST_IF_RE = re.compile(r"\b(IF|ELSIF|WHILE)\s+(TRUE|FALSE|1|0)\s+(THEN|DO)\b", re.I)
_ST_COMMENT_RE = re.compile(r"//.*$|\(\*.*?\*\)|/\*.*?\*/")
_ST_CALL_RE = re.compile(r"\b(?:JSR|FOR)\s*\(\s*([A-Za-z_]\w*)", re.I)


def check_st_line(line, findings, calls):
    text = line["text"]
    P, R, N = line["program"], line["routine"], line["number"]
    code = _ST_COMMENT_RE.sub("", text)

    def add(c, msg, sev=None):
        findings.append(_finding(c, P, R, f"L{N}", msg, text, sev))

    m = _ST_CONST_IF_RE.search(code)
    if m:
        add("ST-CONDICION-CONSTANTE", f"'{m.group(0)}': el bloque siempre o nunca se ejecuta")
    dry = bool(_DRY_RE.search(text))
    for m in _ST_ASSIGN_RE.finditer(code):
        dst, src = m.group(1), m.group(2).strip()
        msg = _valor_fijo_msg(src, dst, dry)
        if msg:
            add("VALOR-FIJO", msg, "BAJA" if dry else None)
        if _INPUT_RE.search(dst):
            add("ESCRITURA-ENTRADA", f"Asignación sobre la entrada '{dst}'")
    for m in _ST_CALL_RE.finditer(code):
        calls[P][m.group(1)].append(f"{R} línea {N}")
    for c in _ST_COMMENT_RE.findall(text):
        m = _SUSP_RE.search(c)
        if m:
            add("COMENTARIO-SOSPECHOSO", f"Menciona '{m.group(0)}': {c.strip()[:140]}")
            break


def analyze_project(proj):
    """Ejecuta todas las reglas. Devuelve (hallazgos crudos, estadistica de instrucciones)."""
    findings = []
    ons_bits = defaultdict(list)     # (programa, bit) -> ubicaciones
    ote_tags = defaultdict(list)     # (programa, tag) -> ubicaciones
    otl_tags, otu_tags = defaultdict(list), defaultdict(list)
    timers = defaultdict(list)
    calls = defaultdict(lambda: defaultdict(list))   # programa -> rutina -> ubicaciones
    instr = Counter()

    for r in proj.rungs:
        series = check_rung(r, findings)
        loc = f"{r['routine']} rung {r['number']}"
        for it in walk(series):
            name, args = it[1], it[2]
            instr[name] += 1
            if not args:
                continue
            key = (r["program"], args[0])
            if name in ("ONS", "OSR", "OSF"):
                ons_bits[key].append(loc)
            elif name == "OTE":
                ote_tags[key].append(loc)
            elif name == "OTL":
                otl_tags[key].append(loc)
            elif name == "OTU":
                otu_tags[key].append(loc)
            elif name in ("TON", "TOF", "RTO"):
                timers[key].append(loc)
            elif name in ("JSR", "FOR"):
                calls[r["program"]][args[0]].append(loc)

    for line in proj.st_lines:
        check_st_line(line, findings, calls)

    def listed(locs):
        extra = f" (+{len(locs) - 4})" if len(locs) > 4 else ""
        return ", ".join(locs[:4]) + extra

    # Bit de one-shot reutilizado en el mismo programa.
    for (prog, bit), locs in ons_bits.items():
        if len(locs) > 1:
            findings.append(_finding("ONS-DUPLICADO", prog, None, None,
                                     f"Bit de one shot '{bit}' usado {len(locs)} veces: {listed(locs)}"))
    # Misma bobina OTE en varios renglones (la ultima gana).
    for (prog, tag), locs in ote_tags.items():
        if len(locs) > 1:
            findings.append(_finding("OTE-DUPLICADA", prog, None, None,
                                     f"OTE de '{tag}' en {len(locs)} renglones: {listed(locs)}"))
    # OTL sin OTU.
    for (prog, tag), locs in otl_tags.items():
        if (prog, tag) not in otu_tags:
            findings.append(_finding("OTL-SIN-OTU", prog, None, None,
                                     f"'{tag}' se enclava con OTL ({listed(locs)}) y no se ve OTU "
                                     "en el mismo programa"))
    # Temporizador reutilizado.
    for (prog, tag), locs in timers.items():
        if len(locs) > 1:
            findings.append(_finding("TIMER-DUPLICADO", prog, None, None,
                                     f"Temporizador '{tag}' usado {len(locs)} veces: {listed(locs)}"))

    # Estructura (solo con export completo del controlador).
    scheduled = proj.scheduled
    for pname, p in proj.programs.items():
        if p["disabled"]:
            findings.append(_finding("PROGRAMA-DESHABILITADO", pname, None, None,
                                     f"El programa '{pname}' está deshabilitado"))
        if not proj.full_export:
            continue
        called = calls.get(pname, {})
        for target, locs in sorted(called.items()):
            if target not in p["routines"]:
                findings.append(_finding("JSR-INEXISTENTE", pname, None, None,
                                         f"Se llama a '{target}' ({listed(locs)}) pero la rutina no existe"))
        if p["main"]:
            for rname in sorted(p["routines"]):
                if rname not in called and rname not in (p["main"], p["fault"]):
                    findings.append(_finding("RUTINA-SIN-LLAMAR", pname, rname, None,
                                             f"La rutina '{rname}' no se llama desde ningún JSR"))
        if proj.tasks and pname not in scheduled and not p["parent"]:
            findings.append(_finding("PROGRAMA-SIN-TAREA", pname, None, None,
                                     f"El programa '{pname}' no está asignado a ninguna tarea"))
    for t in proj.tasks:
        if t["inhibit"]:
            findings.append(_finding("TAREA-INHIBIDA", None, None, None,
                                     f"La tarea '{t['name']}' ({t['type']}) está inhibida",
                                     prog_list=t["programs"]))
    for scope, tag in proj.forces:
        findings.append(_finding("FORZADO", None if scope == "<controller>" else scope, None, None,
                                 f"El tag '{tag}' tiene datos de forzado activos"))
    return findings, instr


def check_project(path):
    """Compatibilidad con la version anterior: (hallazgos, valores, n_renglones)."""
    proj = load_project(path)
    findings, _ = analyze_project(proj)
    return findings, proj.values, len(proj.rungs)


def resolver_valor_debug(values, tag):
    """Intenta leer el valor actual de un bit a partir de los datos L5K del export."""
    m = re.match(r"([A-Za-z0-9_]+)\.BIT\[(\d+)\]\.(\d+)$", tag)
    if m:
        base, idx, bit = m.group(1), int(m.group(2)), int(m.group(3))
    else:
        m2 = re.match(r"([A-Za-z0-9_]+)\.(\d+)$", tag)
        if m2:
            base, idx, bit = m2.group(1), None, int(m2.group(2))
        elif re.fullmatch(r"[A-Za-z0-9_]+", tag):
            base, idx, bit = tag, None, None
        else:
            return None
    for (_scope, name), raw in values.items():
        if name != base:
            continue
        if bit is None:
            return int(raw) if raw in ("0", "1") else None
        nums = re.findall(r"-?\d+", raw)
        if idx is not None and idx < len(nums):
            word = int(nums[idx])
        elif nums:
            word = int(nums[0])
        else:
            return None
        if word < 0:
            word += 1 << 32
        return (word >> bit) & 1
    return None


def _rung_key(v):
    m = re.search(r"\d+", v or "")
    return int(m.group(0)) if m else -1


def finding_id(f):
    raw = "|".join((f["code"], f["prog"], f["rout"], f["rung"], f["msg"]))
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]


def finalize_findings(raw, proj, enabled=None):
    """Agrupa los bits de debug, filtra reglas, quita duplicados y ordena."""
    name = os.path.basename(proj.path)
    out, seen = [], set()
    debug = defaultdict(list)
    for f in raw:
        if f["code"] == "DEBUG-BIT":
            debug[f["tag"]].append(f)
        else:
            out.append(f)
    for tag, items in debug.items():
        val = resolver_valor_debug(proj.values, tag)
        locs = sorted({f"{f['prog']}/{f['rout']} rung {f['rung']}" for f in items}, key=str.lower)
        progs = sorted({f["prog"] for f in items})
        out.append(_finding(
            "DEBUG-BIT", ", ".join(progs)[:80], "-", None,
            f"'{tag}' condiciona lógica en {len(items)} renglón(es); valor actual = "
            f"{'?' if val is None else val}",
            "Ubicaciones: " + "; ".join(locs[:30]) + (" ..." if len(locs) > 30 else ""),
            tag=tag, value=val))
    result = []
    for f in out:
        if enabled is not None and f["code"] not in enabled:
            continue
        f["id"] = finding_id(f)
        if f["id"] in seen:
            continue
        seen.add(f["id"])
        f["file"] = name
        f["path"] = proj.path
        f["title"] = RULES[f["code"]].title
        f.setdefault("status", "Pendiente")
        result.append(f)
    result.sort(key=lambda f: (SEVERITY_ORDER.get(f["sev"], 9), f["code"], f["prog"],
                               f["rout"], _rung_key(f["rung"])))
    return result


def count_by_sev(findings):
    c = Counter(f["sev"] for f in findings)
    return {s: c.get(s, 0) for s in SEVERITIES}


def risk_index(counts, nlines):
    pts = sum(SEV_WEIGHT[s] * counts.get(s, 0) for s in SEVERITIES)
    return pts * 100.0 / max(nlines, 1)


def risk_grade(score):
    for limit, grade in ((2, "A"), (8, "B"), (20, "C"), (45, "D")):
        if score < limit:
            return grade
    return "E"


@dataclass
class AuditResult:
    path: str
    project: Project
    findings: list
    instr: Counter
    elapsed: float

    @property
    def name(self):
        return os.path.basename(self.path)

    @property
    def nrungs(self):
        return len(self.project.rungs)

    @property
    def nst(self):
        return len(self.project.st_lines)

    @property
    def counts(self):
        return count_by_sev(self.findings)

    @property
    def score(self):
        return risk_index(self.counts, self.nrungs + self.nst)

    @property
    def grade(self):
        return risk_grade(self.score)


def audit_file(path, enabled=None):
    t0 = time.perf_counter()
    proj = load_project(path)
    raw, instr = analyze_project(proj)
    findings = finalize_findings(raw, proj, enabled)
    return AuditResult(path, proj, findings, instr, time.perf_counter() - t0)


# ------------------------------------------------------------ exportacion ----
CSV_HEADER = ["archivo", "severidad", "codigo", "regla", "programa", "rutina", "rung",
              "detalle", "estado", "texto_del_renglon"]


def finding_row(f):
    return [f["file"], f["sev"], f["code"], f["title"], f["prog"], f["rout"], f["rung"],
            f["msg"], f.get("status", "Pendiente"), f.get("text", "")[:500]]


def export_csv(dest, results, findings):
    with open(dest, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(CSV_HEADER)
        w.writerows(finding_row(f) for f in findings)


def _summary(results, findings):
    total = count_by_sev(findings)
    lines = sum(r.nrungs + r.nst for r in results)
    score = risk_index(total, lines)
    return total, lines, score, risk_grade(score)


def export_json(dest, results, findings):
    total, lines, score, grade = _summary(results, findings)
    data = {
        "herramienta": APP_NAME, "version": __version__, "autor": __author__,
        "generado": datetime.now().isoformat(timespec="seconds"),
        "resumen": {"archivos": len(results), "renglones": lines, "severidad": total,
                    "indice_riesgo": round(score, 2), "calificacion": grade},
        "archivos": [{
            "archivo": r.name, "ruta": r.path, "controlador": r.project.info,
            "programas": len(r.project.programs), "renglones": r.nrungs, "lineas_st": r.nst,
            "tags": r.project.tag_count, "severidad": r.counts,
            "indice_riesgo": round(r.score, 2), "calificacion": r.grade,
        } for r in results],
        "hallazgos": [{k: f.get(k) for k in ("file", "sev", "code", "title", "prog", "rout",
                                             "rung", "msg", "status", "text", "id")}
                      for f in findings],
    }
    with open(dest, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)


def xlsx_available():
    try:
        import openpyxl  # noqa: F401
        return True
    except Exception:
        return False


def export_xlsx(dest, results, findings):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fills = {"ALTA": "FDE2E2", "MEDIA": "FEF3C7", "BAJA": "EDE9FE", "INFO": "EEF2F7"}
    head_fill, head_font = PatternFill("solid", fgColor="1F2937"), Font(bold=True, color="FFFFFF")
    wb = Workbook()
    wb.properties.creator = __author__
    wb.properties.title = f"{APP_NAME} - informe de auditoría"
    ws = wb.active
    ws.title = "Resumen"
    total, lines, score, grade = _summary(results, findings)
    ws.append([f"{APP_NAME} {__version__} - informe de auditoría"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generado: {datetime.now():%Y-%m-%d %H:%M}   ·   {APP_CREDIT}"])
    ws.append([])
    head = ["Archivo", "Controlador", "Procesador", "Revisión", "Renglones", "Líneas ST",
            "ALTA", "MEDIA", "BAJA", "INFO", "Índice", "Nota"]
    ws.append(head)
    for c in range(1, len(head) + 1):
        ws.cell(row=4, column=c).fill, ws.cell(row=4, column=c).font = head_fill, head_font
    per_file = defaultdict(list)
    for f in findings:
        per_file[f["path"]].append(f)
    for r in results:
        cnt = count_by_sev(per_file[r.path])
        sc = risk_index(cnt, r.nrungs + r.nst)
        info = r.project.info
        ws.append([r.name, info.get("Name", ""), info.get("ProcessorType", ""),
                   f"{info.get('MajorRev', '')}.{info.get('MinorRev', '')}", r.nrungs, r.nst,
                   cnt["ALTA"], cnt["MEDIA"], cnt["BAJA"], cnt["INFO"], round(sc, 1), risk_grade(sc)])
    ws.append([])
    ws.append(["TOTAL", "", "", "", lines, "", total["ALTA"], total["MEDIA"], total["BAJA"],
               total["INFO"], round(score, 1), grade])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for i, w in enumerate((34, 20, 16, 10, 11, 10, 8, 8, 8, 8, 9, 7), 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("Hallazgos")
    ws2.append(["Archivo", "Severidad", "Código", "Regla", "Programa", "Rutina", "Rung", "Detalle",
                "Estado", "Texto del renglón"])
    for c in range(1, len(CSV_HEADER) + 1):
        ws2.cell(row=1, column=c).fill, ws2.cell(row=1, column=c).font = head_fill, head_font
    for f in findings:
        ws2.append(finding_row(f))
        ws2.cell(row=ws2.max_row, column=2).fill = PatternFill("solid", fgColor=fills.get(f["sev"], "FFFFFF"))
    for i, w in enumerate((26, 10, 24, 28, 18, 18, 8, 70, 14, 80), 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for row in ws2.iter_rows(min_row=2):
        row[7].alignment = Alignment(wrap_text=True, vertical="top")
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions
    wb.save(dest)


_HTML_CSS = """
:root{--bg:#f5f7fb;--card:#fff;--text:#1b2330;--muted:#5f6b7c;--border:#e3e7ee;
--alta:#dc2626;--media:#d97706;--baja:#7c3aed;--info:#64748b;--accent:#2563eb}
@media (prefers-color-scheme:dark){:root{--bg:#0f141a;--card:#171d26;--text:#e6edf3;
--muted:#8b98a9;--border:#2a3441;--alta:#f87171;--media:#fbbf24;--baja:#a78bfa;--info:#94a3b8;--accent:#60a5fa}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--text);
font:14px/1.5 "Segoe UI",system-ui,-apple-system,Roboto,sans-serif}
.wrap{max-width:1280px;margin:0 auto;padding:28px 20px 60px}
header{display:flex;align-items:center;gap:14px;margin-bottom:22px}
.logo{width:42px;height:42px;border-radius:10px;background:var(--accent);display:grid;place-items:center;color:#fff;font-weight:700}
h1{font-size:22px;margin:0}h2{font-size:16px;margin:28px 0 10px}.muted{color:var(--muted)}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px}
.kpi{background:var(--card);border:1px solid var(--border);border-radius:12px;padding:14px 16px;border-top:3px solid var(--c,var(--accent))}
.kpi .v{font-size:28px;font-weight:700;color:var(--c,var(--text))}.kpi .l{font-size:12px;color:var(--muted);text-transform:uppercase;letter-spacing:.04em}
table{width:100%;border-collapse:collapse;background:var(--card);border:1px solid var(--border);border-radius:12px;overflow:hidden}
th,td{padding:8px 10px;text-align:left;border-bottom:1px solid var(--border);vertical-align:top}
th{font-size:12px;text-transform:uppercase;letter-spacing:.04em;color:var(--muted);background:color-mix(in srgb,var(--card) 92%,var(--text))}
tr:last-child td{border-bottom:0}.num{text-align:right;font-variant-numeric:tabular-nums}
.b{display:inline-block;padding:2px 8px;border-radius:999px;font-size:11px;font-weight:700;color:#fff}
.ALTA{background:var(--alta)}.MEDIA{background:var(--media)}.BAJA{background:var(--baja)}.INFO{background:var(--info)}
code{font:12px/1.45 Consolas,"Cascadia Mono",monospace;white-space:pre-wrap;word-break:break-all}
details summary{cursor:pointer;color:var(--accent);font-size:12px}
.bar{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin:8px 0 12px}
.bar input{flex:1;min-width:220px;padding:8px 10px;border-radius:8px;border:1px solid var(--border);background:var(--card);color:var(--text)}
.bar label{font-size:13px;display:flex;gap:4px;align-items:center}
.note{margin-top:24px;font-size:12px;color:var(--muted)}
@media print{.bar{display:none}details{display:block}details>summary{display:none}body{background:#fff}}
"""

_HTML_JS = """
const q=document.getElementById('q'),rows=[...document.querySelectorAll('#hall tbody tr')];
function f(){const s=q.value.toLowerCase(),on=[...document.querySelectorAll('.sv:checked')].map(e=>e.value);
let n=0;rows.forEach(r=>{const ok=on.includes(r.dataset.sev)&&(!s||r.textContent.toLowerCase().includes(s));
r.style.display=ok?'':'none';if(ok)n++});document.getElementById('cnt').textContent=n+' visibles'}
q.addEventListener('input',f);document.querySelectorAll('.sv').forEach(e=>e.addEventListener('change',f));f();
"""


def export_html(dest, results, findings):
    e = html.escape
    total, lines, score, grade = _summary(results, findings)
    per_file = defaultdict(list)
    for f in findings:
        per_file[f["path"]].append(f)
    grade_col = {"A": "#16a34a", "B": "#65a30d", "C": "var(--media)", "D": "#ea580c", "E": "var(--alta)"}
    out = [f"<!doctype html><html lang='es'><head><meta charset='utf-8'>"
           f"<meta name='viewport' content='width=device-width,initial-scale=1'>"
           f"<meta name='author' content='{e(__author__)}'>"
           f"<title>Informe de auditoría L5X</title><style>{_HTML_CSS}</style></head><body><div class='wrap'>",
           f"<header><div class='logo'>L5X</div><div><h1>Informe de auditoría L5X</h1>"
           f"<div class='muted'>Generado {datetime.now():%Y-%m-%d %H:%M} &middot; {APP_NAME} {__version__}"
           f" &middot; {len(results)} archivo(s)</div></div></header>",
           "<div class='kpis'>",
           f"<div class='kpi'><div class='l'>Renglones / líneas</div><div class='v'>{lines:,}</div></div>"]
    for s, var in (("ALTA", "--alta"), ("MEDIA", "--media"), ("BAJA", "--baja"), ("INFO", "--info")):
        out.append(f"<div class='kpi' style='--c:var({var})'><div class='l'>Severidad {s}</div>"
                   f"<div class='v'>{total[s]}</div></div>")
    out.append(f"<div class='kpi' style='--c:{grade_col[grade]}'><div class='l'>Índice de riesgo</div>"
               f"<div class='v'>{grade} <span class='muted' style='font-size:14px'>{score:.1f}</span></div></div></div>")

    out.append("<h2>Archivos analizados</h2><table><thead><tr><th>Archivo</th><th>Controlador</th>"
               "<th>Procesador</th><th>Rev.</th><th class='num'>Renglones</th><th class='num'>ALTA</th>"
               "<th class='num'>MEDIA</th><th class='num'>BAJA</th><th class='num'>Índice</th><th>Nota</th>"
               "</tr></thead><tbody>")
    for r in results:
        cnt = count_by_sev(per_file[r.path])
        sc = risk_index(cnt, r.nrungs + r.nst)
        info = r.project.info
        out.append(f"<tr><td>{e(r.name)}</td><td>{e(info.get('Name', '-'))}</td>"
                   f"<td>{e(info.get('ProcessorType', '-'))}</td>"
                   f"<td>{e(info.get('MajorRev', '?'))}.{e(info.get('MinorRev', '?'))}</td>"
                   f"<td class='num'>{r.nrungs + r.nst:,}</td><td class='num'>{cnt['ALTA']}</td>"
                   f"<td class='num'>{cnt['MEDIA']}</td><td class='num'>{cnt['BAJA']}</td>"
                   f"<td class='num'>{sc:.1f}</td><td><b style='color:{grade_col[risk_grade(sc)]}'>"
                   f"{risk_grade(sc)}</b></td></tr>")
    out.append("</tbody></table>")

    by_rule = Counter(f["code"] for f in findings)
    if by_rule:
        out.append("<h2>Hallazgos por regla</h2><table><thead><tr><th>Regla</th><th>Severidad</th>"
                   "<th class='num'>Cantidad</th><th>Qué revisar</th></tr></thead><tbody>")
        for code, n in sorted(by_rule.items(), key=lambda x: (SEVERITY_ORDER[RULES[x[0]].sev], -x[1])):
            rule = RULES[code]
            out.append(f"<tr><td><b>{e(rule.title)}</b><br><span class='muted'>{e(code)}</span></td>"
                       f"<td><span class='b {rule.sev}'>{rule.sev}</span></td><td class='num'>{n}</td>"
                       f"<td>{e(rule.fix)}</td></tr>")
        out.append("</tbody></table>")

    out.append("<h2>Detalle de hallazgos</h2><div class='bar'><input id='q' placeholder='Buscar tag, rutina, "
               "texto...'>")
    for s in SEVERITIES:
        out.append(f"<label><input type='checkbox' class='sv' value='{s}' checked> {s}</label>")
    out.append("<span id='cnt' class='muted'></span></div><table id='hall'><thead><tr><th>Sev.</th>"
               "<th>Regla</th><th>Archivo</th><th>Ubicación</th><th>Detalle</th><th>Estado</th>"
               "</tr></thead><tbody>")
    for f in findings:
        loc = f"{f['prog']} / {f['rout']}" + (f" / rung {f['rung']}" if f["rung"] != "-" else "")
        code = (f"<details><summary>Ver renglón</summary><code>{e(f['text'])}</code></details>"
                if f.get("text") else "")
        out.append(f"<tr data-sev='{f['sev']}'><td><span class='b {f['sev']}'>{f['sev']}</span></td>"
                   f"<td><b>{e(f['title'])}</b><br><span class='muted'>{e(f['code'])}</span></td>"
                   f"<td>{e(f['file'])}</td><td>{e(loc)}</td><td>{e(f['msg'])}{code}</td>"
                   f"<td>{e(f.get('status', 'Pendiente'))}</td></tr>")
    out.append("</tbody></table><p class='note'>Este informe es resultado de un análisis estático. Cada "
               "hallazgo debe confirmarse en el contexto de la máquina antes de tomar acciones. AFI y "
               "OTL sin OTU suelen ser normales.</p>"
               f"<p class='note'>{APP_NAME} {__version__} &middot; {e(APP_CREDIT)} &middot; "
               f"{e(APP_COPYRIGHT)}</p></div>")
    out.append(f"<script>{_HTML_JS}</script></body></html>")
    with open(dest, "w", encoding="utf-8") as fh:
        fh.write("".join(out))


EXPORTERS = {"csv": export_csv, "html": export_html, "json": export_json, "xlsx": export_xlsx}


# ------------------------------------------------------------ proyecto demo --
DEMO_L5X = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<RSLogix5000Content SchemaRevision="1.0" SoftwareRevision="33.01" TargetName="DEMO_Linea3"
 TargetType="Controller" ContainsContext="false" ExportDate="Mon Sep 22 10:15:00 2026">
<Controller Use="Target" Name="DEMO_Linea3" ProcessorType="1756-L83E" MajorRev="33" MinorRev="11"
 LastModifiedDate="Fri Sep 19 16:42:10 2026">
<Tags>
 <Tag Name="gDebug" TagType="Base" DataType="DINT"><Data Format="L5K"><![CDATA[5]]></Data></Tag>
 <Tag Name="Bypass_Puerta" TagType="Base" DataType="BOOL"><Data Format="L5K"><![CDATA[1]]></Data></Tag>
 <Tag Name="Local:2:O" TagType="Base" DataType="AB:1756_DO:O:0"><Data Format="L5K"><![CDATA[[0]]]></Data>
  <ForceData Format="L5K"><![CDATA[[16#0000_0004]]]></ForceData></Tag>
</Tags>
<Programs>
<Program Name="Estacion10" MainRoutineName="MainRoutine" Disabled="false">
<Tags>
 <Tag Name="T_Espera" TagType="Base" DataType="TIMER"><Data Format="L5K"><![CDATA[[0,1000,0]]]></Data></Tag>
</Tags>
<Routines>
<Routine Name="MainRoutine" Type="RLL"><RLLContent>
 <Rung Number="0" Type="N"><Text><![CDATA[JSR(Seguridad,0)JSR(Inspeccion,0)JSR(Calculos,0)JSR(Rutina_Borrada,0);]]></Text></Rung>
</RLLContent></Routine>
<Routine Name="Seguridad" Type="RLL"><RLLContent>
 <Rung Number="0" Type="N"><Comment><![CDATA[Guarda de puerta - puente temporal para pruebas de arranque]]></Comment>
  <Text><![CDATA[[XIC(Puerta_Cerrada) ,XIO(Puerta_Cerrada) ]XIC(Paro_Emergencia_OK)OTE(Permiso_Marcha);]]></Text></Rung>
 <Rung Number="1" Type="N"><Text><![CDATA[XIC(Sensor_Pieza)XIO(Sensor_Pieza)OTE(Pieza_Presente);]]></Text></Rung>
 <Rung Number="2" Type="N"><Text><![CDATA[XIC(gDebug.2)OTE(Permiso_Marcha);]]></Text></Rung>
 <Rung Number="3" Type="N"><Text><![CDATA[XIC(Reset_Fallas)ONS(Ons_Storage)OTL(Falla_Presion);]]></Text></Rung>
 <Rung Number="4" Type="N"><Text><![CDATA[XIC(Arranque)ONS(Ons_Storage)OTE(Pulso_Arranque);]]></Text></Rung>
 <Rung Number="5" Type="N"><Text><![CDATA[AFI()XIC(Permiso_Marcha)OTE(Motor_Transportador);]]></Text></Rung>
 <Rung Number="6" Type="N"><Text><![CDATA[XIC(Bypass_Puerta)OTE(Local:1:I.Data.3);]]></Text></Rung>
</RLLContent></Routine>
<Routine Name="Inspeccion" Type="RLL"><RLLContent>
 <Rung Number="0" Type="N"><Text><![CDATA[XIC(Camara_Lista)MOV(1,Vision.Status_Res);]]></Text></Rung>
 <Rung Number="1" Type="N"><Text><![CDATA[XIC(DryCycle)MOV(1,Vision.Results[0]);]]></Text></Rung>
 <Rung Number="2" Type="N"><Text><![CDATA[XIC(Inicio_Ciclo)TON(T_Espera,?,?);]]></Text></Rung>
 <Rung Number="3" Type="N"><Text><![CDATA[MOV(0,T_Espera.PRE)XIC(Pieza_Presente)TON(T_Espera,?,?);]]></Text></Rung>
 <Rung Number="4" Type="N"><Text><![CDATA[XIC(Fin_Ciclo)JMP(Salto_Fin);]]></Text></Rung>
 <Rung Number="5" Type="N"><Text><![CDATA[XIC(T_Espera.DN)[OTE(Inspeccion_OK) ,OTE(Luz_Verde) ];]]></Text></Rung>
</RLLContent></Routine>
<Routine Name="Calculos" Type="ST"><STContent>
 <Line Number="0"><![CDATA[IF 1 THEN]]></Line>
 <Line Number="1"><![CDATA[    Measure_Altura := 125; // valor provisional mientras se calibra el sensor]]></Line>
 <Line Number="2"><![CDATA[END_IF;]]></Line>
</STContent></Routine>
<Routine Name="Rutina_Vieja" Type="RLL"><RLLContent>
 <Rung Number="0" Type="N"><Text><![CDATA[XIC(Senal_Vieja)OTE(Salida_Vieja);]]></Text></Rung>
</RLLContent></Routine>
</Routines>
</Program>
<Program Name="Pruebas_Viejas" MainRoutineName="MainRoutine" Disabled="true">
<Routines><Routine Name="MainRoutine" Type="RLL"><RLLContent>
 <Rung Number="0" Type="N"><Text><![CDATA[XIC(Debug_Modo)OTE(Salida_Prueba);]]></Text></Rung>
</RLLContent></Routine></Routines>
</Program>
</Programs>
<Tasks>
 <Task Name="MainTask" Type="CONTINUOUS" Priority="10" InhibitTask="false">
  <ScheduledPrograms><ScheduledProgram Name="Estacion10"/></ScheduledPrograms></Task>
 <Task Name="Tarea_Diagnostico" Type="PERIODIC" Rate="100" Priority="12" InhibitTask="true">
  <ScheduledPrograms/></Task>
</Tasks>
</Controller>
</RSLogix5000Content>
"""


def write_demo(folder=None):
    folder = folder or tempfile.gettempdir()
    path = os.path.join(folder, "DEMO_Linea3.L5X")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(DEMO_L5X)
    return path


# ---------------------------------------------------------- configuracion ----
def config_dir():
    base = os.environ.get("APPDATA") or os.path.join(os.path.expanduser("~"), ".config")
    path = os.path.join(base, "L5XAuditor")
    os.makedirs(path, exist_ok=True)
    return path


def load_config():
    try:
        with open(os.path.join(config_dir(), "config.json"), encoding="utf-8") as fh:
            return json.load(fh)
    except Exception:
        return {}


def save_config(cfg):
    try:
        with open(os.path.join(config_dir(), "config.json"), "w", encoding="utf-8") as fh:
            json.dump(cfg, fh, indent=2)
    except Exception:
        pass


def _reviews_file(path):
    key = hashlib.sha1(os.path.abspath(path).lower().encode("utf-8")).hexdigest()[:20]
    folder = os.path.join(config_dir(), "revisiones")
    os.makedirs(folder, exist_ok=True)
    return os.path.join(folder, key + ".json")


def load_reviews(path):
    try:
        with open(_reviews_file(path), encoding="utf-8") as fh:
            return json.load(fh).get("estados", {})
    except Exception:
        return {}


def save_reviews(path, statuses):
    try:
        with open(_reviews_file(path), "w", encoding="utf-8") as fh:
            json.dump({"archivo": os.path.abspath(path), "estados": statuses}, fh, indent=1)
    except Exception:
        pass


# ------------------------------------------------------------------ CLI ------
def cli(argv):
    ap = argparse.ArgumentParser(
        prog="l5x_auditor", description=f"{APP_NAME} {__version__}: auditoría estática de archivos L5X.")
    ap.add_argument("files", nargs="*", help="archivos .L5X a revisar")
    ap.add_argument("--csv", help="guardar reporte CSV")
    ap.add_argument("--html", help="guardar reporte HTML")
    ap.add_argument("--json", help="guardar reporte JSON")
    ap.add_argument("--xlsx", help="guardar reporte Excel (requiere openpyxl)")
    ap.add_argument("--min-sev", choices=SEVERITIES, default="INFO",
                    help="severidad mínima incluida en los reportes")
    ap.add_argument("--solo-altas", action="store_true", help="equivale a --min-sev ALTA")
    ap.add_argument("--disable", action="append", default=[], metavar="CÓDIGO",
                    help="desactivar una regla (se puede repetir)")
    ap.add_argument("--fail-on", choices=("ALTA", "MEDIA", "BAJA"),
                    help="código de salida 2 si hay hallazgos de esta severidad o mayor")
    ap.add_argument("--list-rules", action="store_true", help="listar reglas y salir")
    ap.add_argument("--demo", action="store_true", help="analizar el proyecto de demostración")
    ap.add_argument("--gui", action="store_true", help="abrir la interfaz grafica")
    ap.add_argument("--version", action="version", version=f"{APP_NAME} {__version__} - {APP_CREDIT}")
    a = ap.parse_args(argv)

    if a.list_rules:
        for r in sorted(RULES.values(), key=lambda r: (SEVERITY_ORDER[r.sev], r.code)):
            print(f"{r.sev:6} {r.code:24} {r.title}")
        return 0
    files = list(a.files) + ([write_demo()] if a.demo else [])
    if a.gui or not files:
        return run_gui(files)

    unknown = [c for c in a.disable if c not in RULES]
    if unknown:
        ap.error("reglas desconocidas: " + ", ".join(unknown))
    enabled = set(RULES) - set(a.disable)
    min_sev = "ALTA" if a.solo_altas else a.min_sev
    results, findings, rc = [], [], 0
    for path in files:
        print(f"\nRevisando {os.path.basename(path)} ...")
        try:
            res = audit_file(path, enabled)
        except Exception as exc:
            print(f"  ERROR: {exc}")
            rc = 1
            continue
        results.append(res)
        findings.extend(f for f in res.findings if SEVERITY_ORDER[f["sev"]] <= SEVERITY_ORDER[min_sev])
        info = res.project.info
        print(f"  {info.get('Name', '?')}  {info.get('ProcessorType', '')}  v{info.get('MajorRev', '?')}"
              f"  |  {res.nrungs} renglones, {res.nst} líneas ST  |  {res.elapsed:.2f}s")
        for code, c in Counter(f["code"] for f in res.findings).most_common():
            print(f"    {c:5d}  {RULES[code].sev:5}  {code}")
        cnt = res.counts
        print(f"  ALTA={cnt['ALTA']}  MEDIA={cnt['MEDIA']}  BAJA={cnt['BAJA']}  INFO={cnt['INFO']}"
              f"  |  índice {res.score:.1f} ({res.grade})")
    for kind in ("csv", "html", "json", "xlsx"):
        dest = getattr(a, kind)
        if dest and results:
            EXPORTERS[kind](dest, results, findings)
            print(f"\nReporte {kind.upper()} generado: {dest}")
    if a.fail_on and any(SEVERITY_ORDER[f["sev"]] <= SEVERITY_ORDER[a.fail_on]
                         for r in results for f in r.findings):
        rc = 2
    return rc


# ------------------------------------------------------------------- GUI -----
try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
    from tkinter import font as tkfont
except Exception:      # entorno sin Tk: el motor y la CLI siguen funcionando
    tk = None

_LabelBase = tk.Label if tk else object
_CanvasBase = tk.Canvas if tk else object
_EntryBase = tk.Entry if tk else object

PALETTES = {
    "oscuro": {
        "bg": "#0d1117", "sidebar": "#10151d", "card": "#161c26", "surface": "#0f141b",
        "border": "#263040", "text": "#e6edf3", "muted": "#8b98a9", "faint": "#5b6778",
        "accent": "#4f8cff", "accent_fg": "#ffffff", "accent_hover": "#3b7bf0", "hover": "#1d2532",
        "sel": "#1e3a66", "ALTA": "#f05252", "MEDIA": "#f5a524", "BAJA": "#a78bfa", "INFO": "#8b9bb0",
        "ok": "#22c55e", "code_bg": "#0b0f15", "syn_inst": "#7cb4ff", "syn_num": "#f5a524",
        "syn_br": "#c678dd", "syn_hl": "#4a3a12",
    },
    "claro": {
        "bg": "#f3f5f9", "sidebar": "#ffffff", "card": "#ffffff", "surface": "#f7f9fc",
        "border": "#e0e5ec", "text": "#1b2330", "muted": "#5b6a7d", "faint": "#94a0b0",
        "accent": "#2563eb", "accent_fg": "#ffffff", "accent_hover": "#1d4ed8", "hover": "#eef2f8",
        "sel": "#d6e4ff", "ALTA": "#dc2626", "MEDIA": "#d97706", "BAJA": "#7c3aed", "INFO": "#64748b",
        "ok": "#16a34a", "code_bg": "#f6f8fb", "syn_inst": "#1d4ed8", "syn_num": "#b45309",
        "syn_br": "#9333ea", "syn_hl": "#fff1b8",
    },
}
GRADE_COLORS = {"A": "#22c55e", "B": "#84cc16", "C": "#f5a524", "D": "#f97316", "E": "#f05252"}


def blend(c1, c2, t):
    """Mezcla c1 sobre c2 con opacidad t."""
    a = [int(c1[i:i + 2], 16) for i in (1, 3, 5)]
    b = [int(c2[i:i + 2], 16) for i in (1, 3, 5)]
    return "#" + "".join(f"{round(x * t + y * (1 - t)):02x}" for x, y in zip(a, b))


def pick_font(candidates, fallback):
    try:
        fams = {f.lower() for f in tkfont.families()}
    except Exception:
        return fallback
    for c in candidates:
        if c.lower() in fams:
            return c
    return fallback


class FlatButton(_LabelBase):
    """Boton plano basado en Label (control total de colores en cualquier SO)."""

    def __init__(self, master, app, text, command=None, kind="secondary", bg=None, small=False):
        self.app, self.kind, self.command, self.enabled = app, kind, command, True
        self.base_bg = bg
        super().__init__(master, text=text, cursor="hand2", bd=0,
                         padx=10 if small else 14, pady=4 if small else 7,
                         font=app.F["small_b"] if small else app.F["btn"])
        self.bind("<Enter>", lambda e: self._paint(True))
        self.bind("<Leave>", lambda e: self._paint(False))
        self.bind("<ButtonRelease-1>", self._click)
        self._paint(False)

    def _colors(self):
        P = self.app.P
        if self.kind == "primary":
            return P["accent"], P["accent_fg"], P["accent_hover"], P["accent"]
        if self.kind == "ghost":
            base = self.base_bg or P["bg"]
            return base, P["muted"], P["hover"], base
        return P["card"], P["text"], P["hover"], P["border"]

    def _paint(self, hover):
        bg, fg, hov, border = self._colors()
        if not self.enabled:
            fg = self.app.P["faint"]
            hover = False
            if self.kind == "primary":
                bg = border = blend(self.app.P["accent"], self.app.P["bg"], 0.35)
                fg = blend(self.app.P["accent_fg"], bg, 0.6)
        self.configure(bg=hov if hover else bg, fg=fg, highlightthickness=1,
                       highlightbackground=border, highlightcolor=border,
                       cursor="hand2" if self.enabled else "arrow")

    def _click(self, _e):
        if self.enabled and self.command:
            self.command()

    def set_enabled(self, on):
        self.enabled = bool(on)
        self._paint(False)


class Pill(_LabelBase):
    """Etiqueta conmutables (filtros de severidad / estado)."""

    def __init__(self, master, app, text, color, active=True, command=None):
        self.app, self.color, self.active, self.command = app, color, active, command
        super().__init__(master, text=text, cursor="hand2", padx=10, pady=3, font=app.F["small_b"], bd=0)
        self.bind("<ButtonRelease-1>", self._toggle)
        self.paint()

    def paint(self):
        P = self.app.P
        if self.active:
            self.configure(bg=blend(self.color, P["card"], 0.18), fg=self.color, highlightthickness=1,
                           highlightbackground=blend(self.color, P["card"], 0.55))
        else:
            self.configure(bg=P["card"], fg=P["faint"], highlightthickness=1,
                           highlightbackground=P["border"])

    def _toggle(self, _e):
        self.active = not self.active
        self.paint()
        if self.command:
            self.command()


class SearchEntry(_EntryBase):
    def __init__(self, master, app, placeholder, on_change):
        self.app, self.placeholder, self.on_change, self._job = app, placeholder, on_change, None
        P = app.P
        super().__init__(master, bg=P["surface"], fg=P["text"], insertbackground=P["text"], relief="flat",
                         highlightthickness=1, highlightbackground=P["border"], highlightcolor=P["accent"],
                         font=app.F["body"], bd=6)
        self._ph = False
        self._show_ph()
        self.bind("<FocusIn>", self._focus_in)
        self.bind("<FocusOut>", lambda e: self._show_ph())
        self.bind("<KeyRelease>", self._changed)
        self.bind("<Escape>", lambda e: self.set_value(""))

    def _show_ph(self):
        if not self.get():
            self._ph = True
            self.insert(0, self.placeholder)
            self.configure(fg=self.app.P["faint"])

    def _focus_in(self, _e):
        if self._ph:
            self.delete(0, "end")
            self._ph = False
            self.configure(fg=self.app.P["text"])

    def value(self):
        return "" if self._ph else self.get()

    def set_value(self, v, notify=True):
        self._focus_in(None)
        self.delete(0, "end")
        self.insert(0, v)
        if self.focus_get() is not self:
            self._show_ph()
        if notify:
            self.on_change()

    def _changed(self, _e):
        if self._job:
            self.after_cancel(self._job)
        self._job = self.after(180, self.on_change)


class BarChart(_CanvasBase):
    """Barras horizontales: data = [(etiqueta, valor, color)]."""

    def __init__(self, master, app, data, max_rows=10, row_h=26):
        self.app, self.data, self.row_h = app, data[:max_rows], row_h
        h = max(len(self.data), 1) * row_h + 8
        super().__init__(master, height=h, bg=app.P["card"], highlightthickness=0, bd=0)
        self.bind("<Configure>", lambda e: self.draw())

    def draw(self):
        self.delete("all")
        P, F = self.app.P, self.app.F
        w = self.winfo_width()
        if not self.data:
            self.create_text(w / 2, self.row_h / 2 + 4, text="Sin datos", fill=P["faint"], font=F["small"])
            return
        fnt = tkfont.Font(font=F["small"])
        lw = min(max(fnt.measure(lbl) for lbl, _, _ in self.data) + 14, int(w * 0.45))
        vmax = max(v for _, v, _ in self.data) or 1
        span = max(w - lw - 50, 10)
        for i, (lbl, val, col) in enumerate(self.data):
            y = 4 + i * self.row_h + self.row_h / 2
            self.create_text(lw - 10, y, text=lbl, anchor="e", fill=P["muted"], font=F["small"])
            self.create_rectangle(lw, y - 8, lw + span, y + 8, fill=P["surface"], width=0)
            bw = max(span * val / vmax, 3)
            self.create_rectangle(lw, y - 8, lw + bw, y + 8, fill=col, width=0)
            self.create_text(lw + bw + 8, y, text=f"{val:,}", anchor="w", fill=P["text"], font=F["small_b"])


class Donut(_CanvasBase):
    def __init__(self, master, app, data, size=176, thick=22):
        self.app, self.data, self.size, self.thick = app, data, size, thick
        super().__init__(master, width=size, height=size, bg=app.P["card"], highlightthickness=0, bd=0)
        self.draw()

    def draw(self):
        P, F, s, t = self.app.P, self.app.F, self.size, self.thick
        pad = t / 2 + 2
        box = (pad, pad, s - pad, s - pad)
        total = sum(v for _, v, _ in self.data)
        self.create_oval(*box, outline=P["surface"], width=t)
        start = 90.0
        for _, v, col in self.data:
            if not v:
                continue
            ext = -359.9 * v / total if total else 0
            self.create_arc(*box, start=start, extent=ext, style="arc", outline=col, width=t)
            start += ext
        self.create_text(s / 2, s / 2 - 8, text=f"{total:,}", fill=P["text"], font=F["kpi"])
        self.create_text(s / 2, s / 2 + 20, text="hallazgos", fill=P["muted"], font=F["small"])


class Logo(_CanvasBase):
    def __init__(self, master, app, size=34, bg=None):
        super().__init__(master, width=size, height=size, bg=bg or app.P["sidebar"], highlightthickness=0, bd=0)
        P, s = app.P, size
        r = s * 0.26
        pts = [r, 0, s - r, 0, s, 0, s, r, s, s - r, s, s, s - r, s, r, s, 0, s, 0, s - r, 0, r, 0, 0]
        self.create_polygon(pts, smooth=True, fill=P["accent"], outline="")
        self.create_line(s * .26, s * .52, s * .43, s * .69, s * .75, s * .32, fill="#ffffff",
                         width=max(2, s // 11), capstyle="round", joinstyle="round")


class AuditorApp:
    NAV = (("panel", "Panel", "▦"), ("hallazgos", "Hallazgos", "⚠"),
           ("proyecto", "Proyecto", "☰"), ("reglas", "Reglas", "✔"),
           ("registro", "Registro", "≡"))
    TITLES = {
        "panel": ("Panel", "Resumen del análisis y del riesgo de cada proyecto"),
        "hallazgos": ("Hallazgos", "Filtra, revisa y clasifica cada hallazgo"),
        "proyecto": ("Proyecto", "Estructura del controlador, programas y uso de instrucciones"),
        "reglas": ("Reglas", "Activa o desactiva las reglas de análisis"),
        "registro": ("Registro", "Bitácora de la sesión"),
    }

    def __init__(self, root, files=(), dnd=False):
        self.root, self.dnd = root, dnd
        self.cfg = load_config()
        self.theme = self.cfg.get("theme", "oscuro")
        self.files, self.results, self.findings = [], [], []
        self.reviews = {}
        self.enabled = set(RULES) - set(self.cfg.get("disabled_rules", []))
        self.page = "panel"
        self.flt = {"sev": set(SEVERITIES), "file": 0, "rule": 0, "status": 0, "q": ""}
        self.sort = ("sev", False)
        self.view = []
        self.current = None
        self.log_lines = []
        self.busy, self.q, self.t0 = False, queue.Queue(), 0.0
        self.proj_idx = 0
        self.only_view = tk.BooleanVar(root, value=self.cfg.get("export_only_view", False))
        self._toast = None

        root.title(APP_TITLE)
        root.geometry(self.cfg.get("geometry", "1320x820"))
        root.minsize(1100, 680)
        self._fonts()
        self.build()
        self.add_files(files, quiet=True)
        self.log(f"{APP_NAME} {__version__} listo.")
        self.log("Exporta desde Studio 5000: File > Save As > Logix Designer XML File (*.L5X)")
        root.protocol("WM_DELETE_WINDOW", self.on_close)
        for seq, fn in (("<Control-o>", self.choose_files), ("<F5>", self.analyze),
                        ("<Control-f>", self.focus_search), ("<Control-e>", lambda: self.export("html")),
                        ("<F1>", self.about)):
            root.bind(seq, lambda e, fn=fn: fn())
        if dnd:
            try:
                root.drop_target_register("DND_Files")
                root.dnd_bind("<<Drop>>", lambda e: self.add_files(root.tk.splitlist(e.data)))
            except Exception:
                self.dnd = False
        root.after(60, self._poll)

    # ------------------------------------------------------------ estilo ---
    def _fonts(self):
        ui = pick_font(("Segoe UI", "SF Pro Text", "Helvetica Neue", "Inter", "Ubuntu", "Cantarell",
                        "Noto Sans", "DejaVu Sans"), "TkDefaultFont")
        mono = pick_font(("Cascadia Mono", "Consolas", "JetBrains Mono", "SF Mono", "Menlo",
                          "DejaVu Sans Mono", "Courier New"), "TkFixedFont")
        self.F = {"body": (ui, 10), "small": (ui, 9), "small_b": (ui, 9, "bold"), "btn": (ui, 10, "bold"),
                  "h1": (ui, 17, "bold"), "h2": (ui, 12, "bold"), "h3": (ui, 10, "bold"),
                  "kpi": (ui, 22, "bold"), "logo": (ui, 13, "bold"), "mono": (mono, 10),
                  "caps": (ui, 8, "bold"), "nav": (ui, 10)}

    def _style(self):
        P, F = self.P, self.F
        st = ttk.Style(self.root)
        st.theme_use("clam")
        st.configure(".", background=P["bg"], foreground=P["text"], fieldbackground=P["surface"],
                     bordercolor=P["border"], lightcolor=P["border"], darkcolor=P["border"],
                     troughcolor=P["bg"], focuscolor=P["accent"], selectbackground=P["sel"],
                     selectforeground=P["text"], insertcolor=P["text"], font=F["body"])
        st.configure("Treeview", background=P["card"], fieldbackground=P["card"], foreground=P["text"],
                     rowheight=30, borderwidth=0, relief="flat")
        st.map("Treeview", background=[("selected", P["sel"])], foreground=[("selected", P["text"])])
        st.configure("Treeview.Heading", background=P["card"], foreground=P["muted"], relief="flat",
                     font=F["small_b"], padding=(8, 7), borderwidth=0)
        st.map("Treeview.Heading", background=[("active", P["hover"])])
        st.layout("Treeview", [("Treeview.treearea", {"sticky": "nswe"})])
        st.configure("TCombobox", fieldbackground=P["surface"], background=P["surface"], foreground=P["text"],
                     arrowcolor=P["muted"], padding=5, bordercolor=P["border"])
        st.map("TCombobox", fieldbackground=[("readonly", P["surface"])], foreground=[("readonly", P["text"])],
               selectbackground=[("readonly", P["surface"])], selectforeground=[("readonly", P["text"])],
               bordercolor=[("focus", P["accent"])])
        for opt, val in (("background", P["card"]), ("foreground", P["text"]),
                         ("selectBackground", P["sel"]), ("selectForeground", P["text"])):
            self.root.option_add(f"*TCombobox*Listbox.{opt}", val)
        self.root.option_add("*TCombobox*Listbox.font", F["body"])
        for orient in ("Vertical", "Horizontal"):
            st.configure(f"{orient}.TScrollbar", background=P["border"], troughcolor=P["card"],
                         bordercolor=P["card"], arrowcolor=P["muted"], gripcount=0, relief="flat",
                         lightcolor=P["border"], darkcolor=P["border"], arrowsize=12)
            st.map(f"{orient}.TScrollbar", background=[("active", P["faint"])])
        st.configure("Horizontal.TProgressbar", troughcolor=P["border"], background=P["accent"],
                     bordercolor=P["sidebar"], lightcolor=P["accent"], darkcolor=P["accent"], thickness=6)
        st.configure("TPanedwindow", background=P["bg"])
        st.configure("Sash", sashthickness=10, gripcount=0, background=P["bg"])

    # ------------------------------------------------------- construccion --
    def build(self):
        self.P = PALETTES.get(self.theme, PALETTES["oscuro"])
        P = self.P
        for w in self.root.winfo_children():
            w.destroy()
        self._style()
        self.root.configure(bg=P["bg"])
        self._set_icon()

        side = tk.Frame(self.root, bg=P["sidebar"], width=236)
        side.pack(side="left", fill="y")
        side.pack_propagate(False)
        tk.Frame(self.root, bg=P["border"], width=1).pack(side="left", fill="y")
        self._build_sidebar(side)

        main = tk.Frame(self.root, bg=P["bg"])
        main.pack(side="left", fill="both", expand=True)
        self._build_topbar(main)
        self._build_statusbar(main)
        content = tk.Frame(main, bg=P["bg"])
        content.pack(fill="both", expand=True, padx=24, pady=(4, 16))
        content.rowconfigure(0, weight=1)
        content.columnconfigure(0, weight=1)
        self.pages = {}
        for key, *_ in self.NAV:
            f = tk.Frame(content, bg=P["bg"])
            f.grid(row=0, column=0, sticky="nsew")
            self.pages[key] = f
        self._build_findings_page(self.pages["hallazgos"])
        self._build_rules_page(self.pages["reglas"])
        self._build_log_page(self.pages["registro"])
        self.refresh_all()
        self.show(self.page)

    def _set_icon(self):
        ico, png = resource_path("assets", "l5x_auditor.ico"), resource_path("assets", "l5x_auditor.png")
        try:
            if sys.platform == "win32" and os.path.isfile(ico):
                self.root.iconbitmap(default=ico)
            if os.path.isfile(png):
                self._icon = tk.PhotoImage(master=self.root, file=png)
                self.root.iconphoto(True, self._icon)
                return
        except Exception:
            pass
        try:
            s, P = 32, self.P
            img = tk.PhotoImage(master=self.root, width=s, height=s)
            acc = P["accent"]
            rows = []
            for y in range(s):
                row = []
                for x in range(s):
                    dx, dy = max(6 - x, x - (s - 7), 0), max(6 - y, y - (s - 7), 0)
                    inside = dx * dx + dy * dy <= 36
                    # check mark
                    on = (abs((y - 17) - (x - 8)) <= 2 and 8 <= x <= 14) or \
                         (abs((y - 23) + (x - 14) * 0.9) <= 2 and 13 <= x <= 25)
                    row.append("#ffffff" if inside and on else (acc if inside else P["bg"]))
                rows.append("{" + " ".join(row) + "}")
            img.put(" ".join(rows))
            self.root.iconphoto(True, img)
            self._icon = img
        except Exception:
            pass

    def _build_sidebar(self, side):
        P, F = self.P, self.F
        head = tk.Frame(side, bg=P["sidebar"])
        head.pack(fill="x", padx=18, pady=(20, 18))
        Logo(head, self).pack(side="left")
        tt = tk.Frame(head, bg=P["sidebar"])
        tt.pack(side="left", padx=10)
        tk.Label(tt, text=APP_NAME, font=F["logo"], bg=P["sidebar"], fg=P["text"]).pack(anchor="w")
        tk.Label(tt, text="Studio 5000 · Logix", font=F["small"], bg=P["sidebar"],
                 fg=P["muted"]).pack(anchor="w")

        tk.Label(side, text="NAVEGACIÓN", font=F["caps"], bg=P["sidebar"], fg=P["faint"]).pack(
            anchor="w", padx=20, pady=(4, 4))
        self.nav = {}
        for key, label, icon in self.NAV:
            row = tk.Frame(side, bg=P["sidebar"], cursor="hand2")
            row.pack(fill="x", padx=10, pady=1)
            bar = tk.Frame(row, bg=P["sidebar"], width=3)
            bar.pack(side="left", fill="y")
            ic = tk.Label(row, text=icon, width=2, font=F["body"], bg=P["sidebar"], fg=P["muted"])
            ic.pack(side="left", padx=(10, 4), pady=7)
            lb = tk.Label(row, text=label, font=F["nav"], bg=P["sidebar"], fg=P["muted"], anchor="w")
            lb.pack(side="left", fill="x", expand=True)
            badge = tk.Label(row, text="", font=F["caps"], bg=P["sidebar"], fg=P["sidebar"], padx=6)
            badge.pack(side="right", padx=8)
            self.nav[key] = (row, bar, ic, lb, badge)
            for w in (row, ic, lb, badge):
                w.bind("<ButtonRelease-1>", lambda e, k=key: self.show(k))
                w.bind("<Enter>", lambda e, k=key: self._nav_hover(k, True))
                w.bind("<Leave>", lambda e, k=key: self._nav_hover(k, False))

        tk.Frame(side, bg=P["border"], height=1).pack(fill="x", padx=18, pady=16)
        fh = tk.Frame(side, bg=P["sidebar"])
        fh.pack(fill="x", padx=20)
        self.files_lbl = tk.Label(fh, text="ARCHIVOS", font=F["caps"], bg=P["sidebar"], fg=P["faint"])
        self.files_lbl.pack(side="left")
        FlatButton(fh, self, "✕", self.remove_selected, kind="ghost", bg=P["sidebar"],
                   small=True).pack(side="right")
        FlatButton(fh, self, "+", self.choose_files, kind="ghost", bg=P["sidebar"], small=True).pack(side="right")

        bottom = tk.Frame(side, bg=P["sidebar"])
        bottom.pack(side="bottom", fill="x", padx=18, pady=14)
        credit = [tk.Label(bottom, text="DESARROLLADO POR", font=F["caps"], bg=P["sidebar"], fg=P["faint"]),
                  tk.Label(bottom, text=__author__, font=F["small_b"], bg=P["sidebar"], fg=P["muted"]),
                  tk.Label(bottom, text=f"v{__version__}  ·  Acerca de (F1)", font=F["small"],
                           bg=P["sidebar"], fg=P["faint"])]
        for w in credit:
            w.pack(anchor="w")
            w.configure(cursor="hand2")
            w.bind("<ButtonRelease-1>", lambda e: self.about())
        credit[1].bind("<Enter>", lambda e: credit[1].configure(fg=P["accent"]))
        credit[1].bind("<Leave>", lambda e: credit[1].configure(fg=P["muted"]))

        box = tk.Frame(side, bg=P["sidebar"])
        box.pack(fill="both", expand=True, padx=14, pady=(8, 0))
        self.files_box = tk.Listbox(box, bg=P["sidebar"], fg=P["text"], bd=0, highlightthickness=0,
                                    selectbackground=P["hover"], selectforeground=P["text"],
                                    activestyle="none", font=F["small"], selectmode="extended")
        self.files_box.pack(fill="both", expand=True)
        self.files_box.bind("<Delete>", lambda e: self.remove_selected())
        self.files_hint = tk.Label(box, bg=P["sidebar"], fg=P["faint"], font=F["small"], justify="left",
                                   wraplength=190, anchor="w")

    def _nav_hover(self, key, on):
        if key == self.page:
            return
        row, bar, ic, lb, badge = self.nav[key]
        bg = self.P["hover"] if on else self.P["sidebar"]
        for w in (row, ic, lb):
            w.configure(bg=bg)
        bar.configure(bg=bg)
        if not badge.cget("text"):
            badge.configure(bg=bg, fg=bg)

    def _build_topbar(self, main):
        P, F = self.P, self.F
        top = tk.Frame(main, bg=P["bg"])
        top.pack(fill="x", padx=24, pady=(18, 12))
        left = tk.Frame(top, bg=P["bg"])
        left.pack(side="left")
        self.title_lbl = tk.Label(left, font=F["h1"], bg=P["bg"], fg=P["text"])
        self.title_lbl.pack(anchor="w")
        self.sub_lbl = tk.Label(left, font=F["small"], bg=P["bg"], fg=P["muted"])
        self.sub_lbl.pack(anchor="w")
        right = tk.Frame(top, bg=P["bg"])
        right.pack(side="right")
        icon = "☀  Claro" if self.theme == "oscuro" else "☾  Oscuro"
        FlatButton(right, self, icon, self.toggle_theme, kind="ghost").pack(side="right", padx=(8, 0))
        self.btn_analyze = FlatButton(right, self, "▶  Analizar", self.analyze, kind="primary")
        self.btn_analyze.pack(side="right", padx=(8, 0))
        self.btn_export = FlatButton(right, self, "⤓  Exportar  ▾", self.export_menu)
        self.btn_export.pack(side="right", padx=(8, 0))
        FlatButton(right, self, "+  Agregar L5X", self.choose_files).pack(side="right")

    def _build_statusbar(self, main):
        P, F = self.P, self.F
        bar = tk.Frame(main, bg=P["sidebar"], height=30)
        bar.pack(side="bottom", fill="x")
        tk.Frame(main, bg=P["border"], height=1).pack(side="bottom", fill="x")
        self.status_dot = tk.Label(bar, text="●", font=F["small"], bg=P["sidebar"], fg=P["ok"])
        self.status_dot.pack(side="left", padx=(14, 4), pady=5)
        self.status_lbl = tk.Label(bar, text="Listo", font=F["small"], bg=P["sidebar"], fg=P["muted"])
        self.status_lbl.pack(side="left")
        tk.Label(bar, text="Ctrl+O agregar  ·  F5 analizar  ·  Ctrl+F buscar  ·  Ctrl+E informe HTML",
                 font=F["small"], bg=P["sidebar"], fg=P["faint"]).pack(side="right", padx=14)
        self.progress = ttk.Progressbar(bar, mode="determinate", length=220)

    # ---------------------------------------------------------- utilidades --
    def card(self, master, pad=16, **kw):
        outer = tk.Frame(master, bg=self.P["card"], highlightthickness=1,
                         highlightbackground=self.P["border"], **kw)
        inner = tk.Frame(outer, bg=self.P["card"])
        inner.pack(fill="both", expand=True, padx=pad, pady=pad)
        return outer, inner

    def label(self, master, text="", font="body", fg="text", bg="card", **kw):
        return tk.Label(master, text=text, font=self.F[font], fg=self.P.get(fg, fg),
                        bg=self.P.get(bg, bg), **kw)

    def tree(self, master, columns, heads, widths, height=10, stretch=None, sortable=False):
        P = self.P
        wrap = tk.Frame(master, bg=P["card"])
        tv = ttk.Treeview(wrap, columns=columns, show="headings", height=height, selectmode="browse")
        for c, h, w in zip(columns, heads, widths):
            if sortable:
                tv.heading(c, text=h, anchor="w", command=lambda c=c: self.sort_by(c))
            else:
                tv.heading(c, text=h, anchor="w")
            tv.column(c, width=w, minwidth=40, anchor="w", stretch=(c == stretch))
        vs = ttk.Scrollbar(wrap, orient="vertical", command=tv.yview)
        tv.configure(yscrollcommand=vs.set)
        vs.pack(side="right", fill="y")
        tv.pack(side="left", fill="both", expand=True)
        for s in SEVERITIES:
            tv.tag_configure(s, background=blend(P[s], P["card"], 0.07))
        tv.tag_configure("dim", foreground=P["faint"])
        return wrap, tv

    def log(self, msg):
        line = f"{datetime.now():%H:%M:%S}  {msg}"
        self.log_lines.append(line)
        if getattr(self, "log_txt", None) is not None and self.log_txt.winfo_exists():
            self.log_txt.configure(state="normal")
            self.log_txt.insert("end", line + "\n")
            self.log_txt.see("end")
            self.log_txt.configure(state="disabled")

    def set_status(self, text, color="ok"):
        self.status_lbl.configure(text=text)
        self.status_dot.configure(fg=self.P[color])

    def toast(self, text, color="accent", ms=3800):
        if self._toast is not None:
            try:
                self._toast.destroy()
            except Exception:
                pass
        P = self.P
        t = tk.Frame(self.root, bg=P["card"], highlightthickness=1, highlightbackground=P[color])
        tk.Frame(t, bg=P[color], width=4).pack(side="left", fill="y")
        tk.Label(t, text=text, bg=P["card"], fg=P["text"], font=self.F["body"], padx=14, pady=10,
                 justify="left", wraplength=380).pack(side="left")
        t.place(relx=1.0, rely=1.0, x=-22, y=-46, anchor="se")
        t.lift()
        self._toast = t
        self.root.after(ms, lambda: t.winfo_exists() and t.destroy())

    # ---------------------------------------------------------- navegacion --
    def show(self, key):
        self.page = key
        P = self.P
        for k, (row, bar, ic, lb, badge) in self.nav.items():
            on = k == key
            bg = P["hover"] if on else P["sidebar"]
            for w in (row, ic, lb):
                w.configure(bg=bg)
            bar.configure(bg=P["accent"] if on else bg)
            ic.configure(fg=P["accent"] if on else P["muted"])
            lb.configure(fg=P["text"] if on else P["muted"], font=self.F["h3"] if on else self.F["nav"])
            if not badge.cget("text"):
                badge.configure(bg=bg, fg=bg)
        self.pages[key].tkraise()
        title, sub = self.TITLES[key]
        self.title_lbl.configure(text=title)
        self.sub_lbl.configure(text=sub)

    def focus_search(self):
        self.show("hallazgos")
        self.search.focus_set()

    def toggle_theme(self):
        self.theme = "claro" if self.theme == "oscuro" else "oscuro"
        self.cfg["theme"] = self.theme
        save_config(self.cfg)
        self.build()

    def refresh_all(self):
        self.refresh_files()
        self.render_panel()
        self.refresh_filters()
        self.refresh_findings()
        self.render_project()
        self.refresh_rules()
        self.refresh_badges()

    def refresh_badges(self):
        cnt = count_by_sev(self.findings)
        P = self.P
        items = {"hallazgos": (cnt["ALTA"], "ALTA") if cnt["ALTA"] else (len(self.findings), "INFO")}
        items["reglas"] = (len(RULES) - len(self.enabled), "MEDIA")
        for key, (_row, _bar, _ic, _lb, badge) in self.nav.items():
            n, sev = items.get(key, (0, "INFO"))
            if n:
                badge.configure(text=str(n), bg=blend(P[sev], P["sidebar"], 0.22), fg=P[sev])
            else:
                bg = P["hover"] if key == self.page else P["sidebar"]
                badge.configure(text="", bg=bg, fg=bg)

    # ------------------------------------------------------------ archivos --
    def choose_files(self):
        if self.busy:
            return
        paths = filedialog.askopenfilenames(
            parent=self.root, title="Selecciona archivos L5X", initialdir=self.cfg.get("last_dir") or None,
            filetypes=[("Proyectos Logix (L5X)", "*.L5X *.l5x"), ("Todos los archivos", "*.*")])
        if paths:
            self.cfg["last_dir"] = os.path.dirname(paths[0])
            save_config(self.cfg)
            self.add_files(paths)

    def add_files(self, paths, quiet=False):
        added = 0
        for p in paths:
            p = os.path.abspath(p)
            if p not in self.files and os.path.isfile(p):
                self.files.append(p)
                added += 1
                self.log(f"Archivo agregado: {p}")
        if added:
            self.refresh_files()
            self.render_panel()
            if not quiet:
                self.toast(f"{added} archivo(s) agregado(s). Pulsa Analizar (F5) para revisarlos.")

    def load_demo(self):
        self.add_files([write_demo()])
        self.analyze()

    def remove_selected(self):
        if self.busy:
            return
        sel = list(self.files_box.curselection())
        if not sel:
            return
        gone = {self.files[i] for i in sel}
        self.files = [f for f in self.files if f not in gone]
        self.results = [r for r in self.results if r.path not in gone]
        self.findings = [f for f in self.findings if f["path"] not in gone]
        self.proj_idx = 0
        self.flt["file"] = 0
        self.refresh_all()

    def refresh_files(self):
        self.files_box.delete(0, "end")
        done = {r.path: r for r in self.results}
        for p in self.files:
            r = done.get(p)
            mark = f"  [{r.grade}]" if r else ""
            self.files_box.insert("end", f"  {os.path.basename(p)}{mark}")
        for i, p in enumerate(self.files):
            r = done.get(p)
            if r:
                self.files_box.itemconfigure(i, fg=GRADE_COLORS[r.grade] if r.grade in "DE" else self.P["text"])
        self.files_lbl.configure(text=f"ARCHIVOS  ({len(self.files)})")
        if not self.files:
            self.files_hint.configure(text="Sin archivos. Usa + o Ctrl+O" +
                                      (" o arrastra archivos .L5X a la ventana." if self.dnd else "."))
            self.files_hint.place(x=6, y=4)
        else:
            self.files_hint.place_forget()
        self.btn_analyze.set_enabled(bool(self.files) and not self.busy)
        self.btn_export.set_enabled(bool(self.results) and not self.busy)

    # ------------------------------------------------------------- analisis --
    def analyze(self):
        if self.busy or not self.files:
            if not self.files:
                self.toast("Agrega al menos un archivo .L5X", "MEDIA")
            return
        self.busy = True
        self.t0 = time.perf_counter()
        self.results, self.findings, self.current = [], [], None
        self.btn_analyze.set_enabled(False)
        self.btn_export.set_enabled(False)
        self.progress.configure(maximum=len(self.files), value=0)
        self.progress.pack(side="right", padx=10)
        self.set_status("Analizando...", "MEDIA")
        self.log(f"Iniciando análisis de {len(self.files)} archivo(s) con {len(self.enabled)} reglas activas")
        threading.Thread(target=self._worker, args=(list(self.files), set(self.enabled)), daemon=True).start()

    def _worker(self, files, enabled):
        for i, path in enumerate(files):
            self.q.put(("progress", i, os.path.basename(path)))
            try:
                self.q.put(("result", audit_file(path, enabled)))
            except Exception as exc:
                self.q.put(("error", path, exc))
        self.q.put(("done",))

    def _poll(self):
        try:
            while True:
                msg = self.q.get_nowait()
                kind = msg[0]
                if kind == "progress":
                    self.progress.configure(value=msg[1])
                    self.set_status(f"Analizando {msg[2]}  ({msg[1] + 1}/{len(self.files)})", "MEDIA")
                elif kind == "result":
                    self._take_result(msg[1])
                elif kind == "error":
                    self.log(f"ERROR en {os.path.basename(msg[1])}: {msg[2]}")
                    self.toast(f"No se pudo leer {os.path.basename(msg[1])}:\n{msg[2]}", "ALTA", 6000)
                elif kind == "done":
                    self._finish()
        except queue.Empty:
            pass
        self.root.after(60, self._poll)

    def _take_result(self, res):
        rev = load_reviews(res.path)
        self.reviews[res.path] = rev
        for f in res.findings:
            f["status"] = rev.get(f["id"], "Pendiente")
        self.results.append(res)
        self.findings.extend(res.findings)
        cnt = res.counts
        self.log(f"{res.name}: {res.nrungs} renglones, {res.nst} líneas ST en {res.elapsed:.2f}s  ->  "
                 f"ALTA {cnt['ALTA']}, MEDIA {cnt['MEDIA']}, BAJA {cnt['BAJA']}, INFO {cnt['INFO']}  "
                 f"(índice {res.score:.1f}, nota {res.grade})")
        for code, c in Counter(f["code"] for f in res.findings).most_common():
            self.log(f"      {c:5d}  {code}")
        self.progress.configure(value=len(self.results))

    def _finish(self):
        self.busy = False
        self.progress.pack_forget()
        dt = time.perf_counter() - self.t0
        cnt = count_by_sev(self.findings)
        self.findings.sort(key=lambda f: (SEVERITY_ORDER[f["sev"]], f["file"], f["code"], f["prog"],
                                          f["rout"], _rung_key(f["rung"])))
        self.flt["file"] = min(self.flt["file"], len(self.results))
        self.proj_idx = 0
        self.refresh_all()
        self.show("panel")
        msg = f"{len(self.findings)} hallazgos en {len(self.results)} archivo(s) ({dt:.1f}s)"
        self.set_status(f"Análisis completo: {msg}", "ALTA" if cnt["ALTA"] else "ok")
        self.log(f"Análisis completo: {msg}")
        if self.results:
            self.toast(f"Análisis completo\n{cnt['ALTA']} ALTA  ·  {cnt['MEDIA']} MEDIA  ·  "
                       f"{cnt['BAJA']} BAJA", "ALTA" if cnt["ALTA"] else "ok")

    # --------------------------------------------------------------- panel --
    def render_panel(self):
        pg = self.pages["panel"]
        for w in pg.winfo_children():
            w.destroy()
        if not self.results:
            self._render_empty(pg)
            return
        P, F = self.P, self.F
        cnt = count_by_sev(self.findings)
        lines = sum(r.nrungs + r.nst for r in self.results)
        score = risk_index(cnt, lines)
        grade = risk_grade(score)

        kpis = tk.Frame(pg, bg=P["bg"])
        kpis.pack(fill="x")
        items = [
            ("RENGLONES ANALIZADOS", f"{lines:,}", f"{len(self.results)} archivo(s)", P["accent"], None),
            ("SEVERIDAD ALTA", str(cnt["ALTA"]), "revisar de inmediato", P["ALTA"], "ALTA"),
            ("SEVERIDAD MEDIA", str(cnt["MEDIA"]), "revisar y documentar", P["MEDIA"], "MEDIA"),
            ("SEVERIDAD BAJA", str(cnt["BAJA"]), f"+ {cnt['INFO']} informativos", P["BAJA"], "BAJA"),
            ("ÍNDICE DE RIESGO", grade, f"{score:.1f} pts / 100 renglones", GRADE_COLORS[grade], None),
        ]
        for i, (title, val, sub, col, sev) in enumerate(items):
            outer = tk.Frame(kpis, bg=P["card"], highlightthickness=1, highlightbackground=P["border"])
            outer.grid(row=0, column=i, sticky="nsew", padx=(0 if i == 0 else 12, 0))
            kpis.columnconfigure(i, weight=1, uniform="k")
            tk.Frame(outer, bg=col, height=3).pack(fill="x")
            inner = tk.Frame(outer, bg=P["card"])
            inner.pack(fill="both", expand=True, padx=16, pady=(10, 12))
            ws = [self.label(inner, title, "caps", "muted"),
                  self.label(inner, val, "kpi", col if sev or i == 4 else "text"),
                  self.label(inner, sub, "small", "faint")]
            for w in ws:
                w.pack(anchor="w")
            if sev:
                for w in [outer, inner] + ws:
                    w.configure(cursor="hand2")
                    w.bind("<ButtonRelease-1>", lambda e, s=sev: self.jump_to_sev(s))

        mid = tk.Frame(pg, bg=P["bg"])
        mid.pack(fill="x", pady=14)
        o1, c1 = self.card(mid, width=420)
        o1.pack(side="left", fill="y")
        self.label(c1, "Distribución por severidad", "h2").pack(anchor="w")
        body = tk.Frame(c1, bg=P["card"])
        body.pack(fill="both", expand=True, pady=(10, 0))
        Donut(body, self, [(s, cnt[s], P[s]) for s in SEVERITIES]).pack(side="left")
        leg = tk.Frame(body, bg=P["card"])
        leg.pack(side="left", padx=(20, 4), fill="y")
        total = sum(cnt.values()) or 1
        for s in SEVERITIES:
            r = tk.Frame(leg, bg=P["card"])
            r.pack(anchor="w", pady=5)
            tk.Frame(r, bg=P[s], width=10, height=10).pack(side="left", padx=(0, 8))
            self.label(r, f"{s:<6}", "small_b").pack(side="left")
            self.label(r, f"{cnt[s]:>5}   {100 * cnt[s] / total:4.0f}%", "small", "muted").pack(side="left", padx=8)

        o2, c2 = self.card(mid)
        o2.pack(side="left", fill="both", expand=True, padx=(12, 0))
        hdr = tk.Frame(c2, bg=P["card"])
        hdr.pack(fill="x")
        self.label(hdr, "Hallazgos por regla", "h2").pack(side="left")
        self.label(hdr, "clic en una barra para filtrar", "small", "faint").pack(side="right")
        by_rule = Counter(f["code"] for f in self.findings)
        data = sorted(((c, n, P[RULES[c].sev]) for c, n in by_rule.items()),
                      key=lambda x: (-x[1], SEVERITY_ORDER[RULES[x[0]].sev]))
        chart = BarChart(c2, self, data, max_rows=7)
        chart.pack(fill="both", expand=True, pady=(10, 0))
        chart.bind("<ButtonRelease-1>", lambda e, d=data: self._chart_click(e, d, chart))
        chart.configure(cursor="hand2")

        o3, c3 = self.card(pg)
        o3.pack(fill="both", expand=True)
        self.label(c3, "Archivos analizados", "h2").pack(anchor="w", pady=(0, 8))
        cols = ("file", "ctrl", "cpu", "rev", "lines", "alta", "media", "baja", "idx", "grade", "t")
        wrap, tv = self.tree(c3, cols, ("Archivo", "Controlador", "Procesador", "Rev.", "Renglones", "ALTA",
                                        "MEDIA", "BAJA", "Índice", "Nota", "Tiempo"),
                             (230, 150, 120, 60, 90, 60, 60, 60, 70, 50, 70), height=4, stretch="file")
        wrap.pack(fill="both", expand=True)
        for i, r in enumerate(self.results):
            c = r.counts
            info = r.project.info
            tv.insert("", "end", iid=str(i), values=(
                r.name, info.get("Name", "-"), info.get("ProcessorType", "-"),
                f"{info.get('MajorRev', '?')}.{info.get('MinorRev', '?')}", f"{r.nrungs + r.nst:,}",
                c["ALTA"], c["MEDIA"], c["BAJA"], f"{r.score:.1f}", r.grade, f"{r.elapsed:.2f}s"),
                tags=("ALTA",) if r.grade in "DE" else ())
        tv.bind("<Double-1>", lambda e: self._open_file_findings(tv))

    def _chart_click(self, e, data, chart):
        i = int((e.y - 4) // chart.row_h)
        if 0 <= i < len(chart.data):
            codes = self._rule_codes()
            code = chart.data[i][0]
            self.flt.update(rule=codes.index(code) + 1 if code in codes else 0, sev=set(SEVERITIES))
            self.refresh_filters()
            self.refresh_findings()
            self.show("hallazgos")

    def _open_file_findings(self, tv):
        sel = tv.selection()
        if sel:
            self.flt["file"] = int(sel[0]) + 1
            self.refresh_filters()
            self.refresh_findings()
            self.show("hallazgos")

    def jump_to_sev(self, sev):
        self.flt.update(sev={sev}, rule=0, status=0)
        self.refresh_filters()
        self.refresh_findings()
        self.show("hallazgos")

    def _render_empty(self, pg):
        P, F = self.P, self.F
        wrap = tk.Frame(pg, bg=P["bg"])
        wrap.place(relx=0.5, rely=0.46, anchor="center")
        outer, c = self.card(wrap, pad=36)
        outer.pack()
        Logo(c, self, size=58, bg=P["card"]).pack()
        self.label(c, "Audita tus proyectos Studio 5000", "h1").pack(pady=(16, 4))
        n = len(self.files)
        sub = (f"{n} archivo(s) listo(s) para analizar." if n else
               "Exporta el proyecto como .L5X (File > Save As > Logix Designer XML File)\n"
               "y agrégalo aquí para detectar lógica puenteada, forzada o mal escrita.")
        self.label(c, sub, "body", "muted", justify="center").pack()
        btns = tk.Frame(c, bg=P["card"])
        btns.pack(pady=(20, 22))
        if n:
            FlatButton(btns, self, "▶  Analizar ahora", self.analyze, kind="primary").pack(side="left", padx=5)
            FlatButton(btns, self, "+  Agregar más", self.choose_files).pack(side="left", padx=5)
        else:
            FlatButton(btns, self, "+  Agregar archivos L5X", self.choose_files, kind="primary").pack(
                side="left", padx=5)
            FlatButton(btns, self, "Probar con proyecto demo", self.load_demo).pack(side="left", padx=5)
        grid = tk.Frame(c, bg=P["card"])
        grid.pack()
        feats = [("Interlocks puenteados", "XIC/XIO en ramas paralelas"),
                 ("Lógica muerta", "AFI, series imposibles, rutinas sin JSR"),
                 ("Resultados forzados", "constantes en inspección/medición"),
                 ("E/S intervenida", "forces y escritura sobre entradas"),
                 ("Errores clasicos", "doble bobina, ONS y timers repetidos"),
                 ("Informes", "HTML, Excel, CSV y JSON")]
        for i, (t, d) in enumerate(feats):
            cell = tk.Frame(grid, bg=P["card"])
            cell.grid(row=i // 3, column=i % 3, sticky="w", padx=14, pady=6)
            self.label(cell, "✔  " + t, "h3", "text").pack(anchor="w")
            self.label(cell, "     " + d, "small", "muted").pack(anchor="w")
        if self.dnd:
            self.label(c, "También puedes arrastrar archivos .L5X a la ventana", "small", "faint").pack(pady=(16, 0))

    # ----------------------------------------------------------- hallazgos --
    def _rule_codes(self):
        return sorted({f["code"] for f in self.findings},
                      key=lambda c: (SEVERITY_ORDER[RULES[c].sev], c))

    def _build_findings_page(self, pg):
        P, F = self.P, self.F
        fo, fb = self.card(pg, pad=12)
        fo.pack(fill="x")
        row1 = tk.Frame(fb, bg=P["card"])
        row1.pack(fill="x")
        self.search = SearchEntry(row1, self, "⌕  Buscar tag, rutina, texto...  (Ctrl+F)", self._filters_changed)
        self.search.pack(side="left", fill="x", expand=True)
        if self.flt["q"]:
            self.search.set_value(self.flt["q"], notify=False)
        pills = tk.Frame(row1, bg=P["card"])
        pills.pack(side="left", padx=(12, 0))
        self.pills = {}
        for s in SEVERITIES:
            p = Pill(pills, self, s, P[s], s in self.flt["sev"], self._filters_changed)
            p.pack(side="left", padx=2)
            self.pills[s] = p
        row2 = tk.Frame(fb, bg=P["card"])
        row2.pack(fill="x", pady=(10, 0))
        self.cb_file = ttk.Combobox(row2, state="readonly", width=30)
        self.cb_rule = ttk.Combobox(row2, state="readonly", width=28)
        self.cb_status = ttk.Combobox(row2, state="readonly", width=18)
        for lbl, cb in (("Archivo", self.cb_file), ("Regla", self.cb_rule), ("Estado", self.cb_status)):
            self.label(row2, lbl, "small", "muted").pack(side="left", padx=(0, 6))
            cb.pack(side="left", padx=(0, 16))
            cb.bind("<<ComboboxSelected>>", lambda e: self._filters_changed())
        FlatButton(row2, self, "✕  Limpiar filtros", self.clear_filters, small=True).pack(side="right")

        info = tk.Frame(pg, bg=P["bg"])
        info.pack(fill="x", pady=(10, 6))
        self.count_lbl = self.label(info, "", "small", "muted", "bg")
        self.count_lbl.pack(side="left")
        self.label(info, "Estado: 1 Pendiente · 2 Confirmado · 3 Falso positivo · 4 Aceptado"
                   "   |   clic derecho para más opciones", "small", "faint", "bg").pack(side="right")

        pane = ttk.Panedwindow(pg, orient="vertical")
        pane.pack(fill="both", expand=True)
        to, tb = self.card(pane, pad=1)
        cols = ("sev", "status", "code", "file", "prog", "rout", "rung", "msg")
        self.heads = dict(zip(cols, ("Severidad", "Estado", "Regla", "Archivo", "Programa", "Rutina",
                                     "Rung", "Detalle")))
        wrap, self.tv = self.tree(tb, cols, [self.heads[c] for c in cols],
                                  (104, 116, 200, 160, 130, 130, 64, 520), height=12, stretch="msg",
                                  sortable=True)
        wrap.pack(fill="both", expand=True)
        self.tv.bind("<<TreeviewSelect>>", lambda e: self._on_select())
        self.tv.bind("<Button-3>", self._context_menu)
        self.tv.bind("<Control-c>", lambda e: self.copy_text())
        for k, st in zip("1234", STATUSES):
            self.tv.bind(k, lambda e, st=st: self.set_status_sel(st))
        pane.add(to, weight=3)

        do, db = self.card(pane, pad=16)
        pane.add(do, weight=2)
        self._build_detail(db)
        self.root.after(200, lambda: pane.winfo_exists() and pane.sashpos(0, max(int(pane.winfo_height() * 0.52), 200)))

    def _build_detail(self, db):
        P, F = self.P, self.F
        head = tk.Frame(db, bg=P["card"])
        head.pack(fill="x")
        self.d_badge = tk.Label(head, font=F["caps"], padx=8, pady=2, bg=P["card"])
        self.d_badge.pack(side="left")
        self.d_title = self.label(head, "Selecciona un hallazgo para ver el detalle", "h2")
        self.d_title.pack(side="left", padx=10)
        acts = tk.Frame(head, bg=P["card"])
        acts.pack(side="right")
        self.d_status = {}
        colors = {"Pendiente": P["muted"], "Confirmado": P["ALTA"], "Falso positivo": P["INFO"],
                  "Aceptado": P["ok"]}
        for st in STATUSES:
            p = Pill(acts, self, st, colors[st], False, None)
            p.bind("<ButtonRelease-1>", lambda e, st=st: self.set_status_sel(st), add=False)
            p.pack(side="left", padx=2)
            self.d_status[st] = p
        FlatButton(acts, self, "⧉  Copiar", self.copy_text, small=True).pack(side="left", padx=(10, 0))

        self.d_loc = self.label(db, "", "small", "muted")
        self.d_loc.pack(anchor="w", pady=(6, 10))
        body = tk.Frame(db, bg=P["card"])
        body.pack(fill="both", expand=True)
        body.columnconfigure(0, weight=2, uniform="d")
        body.columnconfigure(1, weight=3, uniform="d")
        body.rowconfigure(0, weight=1)
        self.d_info = tk.Text(body, bg=P["card"], fg=P["text"], font=F["body"], relief="flat", wrap="word",
                              highlightthickness=0, bd=0, height=6, cursor="arrow", padx=0, pady=0,
                              selectbackground=P["sel"], spacing1=1, spacing3=1)
        self.d_info.grid(row=0, column=0, sticky="nsew", padx=(0, 18))
        self.d_info.tag_configure("h", font=F["caps"], foreground=P["faint"], spacing1=4, spacing3=3)
        self.d_info.tag_configure("msg", font=F["h3"], foreground=P["text"], spacing3=8)
        self.d_info.tag_configure("txt", foreground=P["muted"], spacing3=8)
        self.d_info.configure(state="disabled")
        right = tk.Frame(body, bg=P["card"])
        right.grid(row=0, column=1, sticky="nsew")
        self.d_code_h = self.label(right, "", "caps", "faint")
        self.d_code_h.pack(anchor="w")
        tw = tk.Frame(right, bg=P["border"])
        tw.pack(fill="both", expand=True, pady=(2, 0))
        self.d_code = tk.Text(tw, bg=P["code_bg"], fg=P["text"], font=F["mono"], relief="flat", wrap="word",
                              padx=12, pady=10, highlightthickness=0, bd=0, insertbackground=P["text"],
                              height=6, selectbackground=P["sel"])
        self.d_code.pack(fill="both", expand=True, padx=1, pady=1)
        for tag, col in (("inst", P["syn_inst"]), ("num", P["syn_num"]), ("br", P["syn_br"])):
            self.d_code.tag_configure(tag, foreground=col)
        self.d_code.tag_configure("inst", font=F["mono"] + ("bold",))
        self.d_code.tag_configure("hl", background=P["syn_hl"])
        self.d_code.tag_configure("cmt", foreground=P["muted"], font=(F["body"][0], 9, "italic"))
        self.d_code.configure(state="disabled")

    def _filters_changed(self):
        self.flt["sev"] = {s for s, p in self.pills.items() if p.active}
        self.flt["file"] = max(self.cb_file.current(), 0)
        self.flt["rule"] = max(self.cb_rule.current(), 0)
        self.flt["status"] = max(self.cb_status.current(), 0)
        self.flt["q"] = self.search.value()
        self.refresh_findings()

    def clear_filters(self):
        self.flt.update(sev=set(SEVERITIES), file=0, rule=0, status=0, q="")
        self.search.set_value("", notify=False)
        self.refresh_filters()
        self.refresh_findings()

    def refresh_filters(self):
        self.cb_file.configure(values=["Todos los archivos"] + [r.name for r in self.results])
        self.cb_rule.configure(values=["Todas las reglas"] + self._rule_codes())
        self.cb_status.configure(values=["Todos los estados"] + list(STATUSES))
        for cb, key in ((self.cb_file, "file"), (self.cb_rule, "rule"), (self.cb_status, "status")):
            n = len(cb.cget("values"))
            self.flt[key] = self.flt[key] if self.flt[key] < n else 0
            cb.current(self.flt[key])
        for s, p in self.pills.items():
            p.active = s in self.flt["sev"]
            p.paint()

    def _visible(self):
        fl = self.flt
        q = fl["q"].lower().strip()
        path = self.results[fl["file"] - 1].path if fl["file"] and fl["file"] <= len(self.results) else None
        codes = self._rule_codes()
        code = codes[fl["rule"] - 1] if fl["rule"] and fl["rule"] <= len(codes) else None
        status = STATUSES[fl["status"] - 1] if fl["status"] else None
        out = []
        for i, f in enumerate(self.findings):
            if f["sev"] not in fl["sev"] or (path and f["path"] != path) or (code and f["code"] != code):
                continue
            if status and f["status"] != status:
                continue
            if q and q not in " ".join((f["msg"], f["text"], f["prog"], f["rout"], f["code"],
                                        f["file"], f["title"])).lower():
                continue
            out.append(i)
        col, rev = self.sort
        keyf = {"sev": lambda f: (SEVERITY_ORDER[f["sev"]], f["code"]),
                "rung": lambda f: (_rung_key(f["rung"]), f["prog"]),
                "status": lambda f: STATUSES.index(f["status"])}.get(col, lambda f: str(f.get(col, "")).lower())
        out.sort(key=lambda i: keyf(self.findings[i]), reverse=rev)
        return out

    def sort_by(self, col):
        c, rev = self.sort
        self.sort = (col, not rev if c == col else False)
        self.refresh_findings()

    def _row(self, f):
        return (f["sev"], f["status"], f["code"], f["file"], f["prog"], f["rout"],
                f["rung"], f["msg"])

    def _tags(self, f):
        return (f["sev"], "dim") if f["status"] in ("Falso positivo", "Aceptado") else (f["sev"],)

    def refresh_findings(self):
        tv = self.tv
        tv.delete(*tv.get_children())
        self.view = self._visible()
        for i in self.view:
            f = self.findings[i]
            tv.insert("", "end", iid=str(i), values=self._row(f), tags=self._tags(f))
        col, rev = self.sort
        for c, h in self.heads.items():
            tv.heading(c, text=h + ((" ▼" if rev else " ▲") if c == col else ""))
        n, tot = len(self.view), len(self.findings)
        self.count_lbl.configure(text=f"Mostrando {n:,} de {tot:,} hallazgos" if tot else
                                 "Aún no hay resultados. Agrega archivos y pulsa Analizar.")
        if self.current is not None and str(self.current) in tv.get_children():
            tv.selection_set(str(self.current))
            tv.see(str(self.current))
        else:
            self.show_detail(None)

    def _on_select(self):
        sel = self.tv.selection()
        self.current = int(sel[0]) if sel else None
        self.show_detail(self.findings[self.current] if sel else None)

    def show_detail(self, f):
        P = self.P
        for st, p in self.d_status.items():
            p.active = bool(f) and f["status"] == st
            p.paint()
        code, info = self.d_code, self.d_info
        for w in (code, info):
            w.configure(state="normal")
            w.delete("1.0", "end")
        if not f:
            self.d_badge.configure(text="", bg=P["card"])
            self.d_title.configure(text="Selecciona un hallazgo para ver el detalle")
            for w in (self.d_loc, self.d_code_h):
                w.configure(text="")
            for w in (code, info):
                w.configure(state="disabled")
            return
        rule = RULES[f["code"]]
        self.d_badge.configure(text=f["sev"], bg=P[f["sev"]], fg="#ffffff")
        self.d_title.configure(text=rule.title)
        loc = f"{f['file']}   ›   {f['prog']}   ›   {f['rout']}"
        if f["rung"] != "-":
            loc += f"   ›   {'línea ' + f['rung'][1:] if f['rung'].startswith('L') else 'rung ' + f['rung']}"
        self.d_loc.configure(text=f"{loc}      [{f['code']} · {rule.category}]")
        for head, body, tag in (("HALLAZGO", f["msg"], "msg"), ("POR QUÉ IMPORTA", rule.desc, "txt"),
                                ("QUÉ REVISAR", rule.fix, "txt")):
            info.insert("end", head + "\n", "h")
            info.insert("end", body + "\n", tag)
        info.configure(state="disabled")
        self.d_code_h.configure(text="TEXTO DEL RENGLÓN" if f.get("text") else "")
        text = f.get("text") or ""
        if text:
            code.insert("1.0", text)
            for pat, tag in ((r"\b[A-Za-z_]\w*(?=\s*\()", "inst"), (r"(?<![\w.#:])-?\d+(\.\d+)?\b(?!:)", "num"),
                             (r"[\[\],]", "br")):
                for m in re.finditer(pat, text):
                    code.tag_add(tag, f"1.0+{m.start()}c", f"1.0+{m.end()}c")
            for name in re.findall(r"'([^']+)'", f["msg"]):
                for m in re.finditer(re.escape(name), text):
                    code.tag_add("hl", f"1.0+{m.start()}c", f"1.0+{m.end()}c")
        else:
            code.insert("1.0", "Hallazgo a nivel de programa / proyecto: no hay un renglón asociado.", "cmt")
        code.configure(state="disabled")

    def _context_menu(self, e):
        iid = self.tv.identify_row(e.y)
        if not iid:
            return
        self.tv.selection_set(iid)
        m = self.menu()
        for k, st in zip("1234", STATUSES):
            m.add_command(label=f"Marcar como {st}", accelerator=k, command=lambda st=st: self.set_status_sel(st))
        m.add_separator()
        m.add_command(label="Copiar hallazgo", accelerator="Ctrl+C", command=self.copy_text)
        m.add_command(label="Filtrar por esta regla", command=self.filter_current_rule)
        m.add_command(label="Ocultar esta regla (desactivar)", command=self.disable_current_rule)
        m.tk_popup(e.x_root, e.y_root)

    def menu(self):
        P = self.P
        return tk.Menu(self.root, tearoff=0, bg=P["card"], fg=P["text"], activebackground=P["sel"],
                       activeforeground=P["text"], bd=1, relief="flat", font=self.F["body"],
                       disabledforeground=P["faint"])

    def filter_current_rule(self):
        if self.current is None:
            return
        codes = self._rule_codes()
        self.flt["rule"] = codes.index(self.findings[self.current]["code"]) + 1
        self.refresh_filters()
        self.refresh_findings()

    def disable_current_rule(self):
        if self.current is None:
            return
        code = self.findings[self.current]["code"]
        self.toggle_rule(code, False)
        self.findings = [f for f in self.findings if f["code"] != code]
        for r in self.results:
            r.findings = [f for f in r.findings if f["code"] != code]
        self.current = None
        self.refresh_all()
        self.toast(f"Regla {code} desactivada. Puedes reactivarla en Reglas.")

    def set_status_sel(self, status):
        if self.current is None:
            return
        f = self.findings[self.current]
        f["status"] = status
        rev = self.reviews.setdefault(f["path"], {})
        if status == "Pendiente":
            rev.pop(f["id"], None)
        else:
            rev[f["id"]] = status
        save_reviews(f["path"], rev)
        iid = str(self.current)
        if self.tv.exists(iid):
            self.tv.item(iid, values=self._row(f), tags=self._tags(f))
        self.show_detail(f)
        nxt = self.tv.next(iid) if self.tv.exists(iid) else ""
        if nxt:
            self.tv.selection_set(nxt)
            self.tv.see(nxt)
            self.tv.focus(nxt)

    def copy_text(self):
        if self.current is None:
            return
        f = self.findings[self.current]
        txt = (f"[{f['sev']}] {f['code']} - {f['title']}\n{f['file']} / {f['prog']} / {f['rout']} / "
               f"rung {f['rung']}\n{f['msg']}\n{f.get('text', '')}")
        self.root.clipboard_clear()
        self.root.clipboard_append(txt)
        self.toast("Hallazgo copiado al portapapeles", ms=1800)

    # ------------------------------------------------------------ proyecto --
    def render_project(self):
        pg = self.pages["proyecto"]
        for w in pg.winfo_children():
            w.destroy()
        P, F = self.P, self.F
        if not self.results:
            self.label(pg, "Analiza un proyecto para ver su estructura.", "body", "muted", "bg").pack(pady=40)
            return
        self.proj_idx = min(self.proj_idx, len(self.results) - 1)
        r = self.results[self.proj_idx]
        proj, info = r.project, r.project.info

        top = tk.Frame(pg, bg=P["bg"])
        top.pack(fill="x", pady=(0, 12))
        self.label(top, "Archivo:", "body", "muted", "bg").pack(side="left")
        cb = ttk.Combobox(top, state="readonly", width=40, values=[x.name for x in self.results])
        cb.current(self.proj_idx)
        cb.pack(side="left", padx=8)
        cb.bind("<<ComboboxSelected>>", lambda e: (setattr(self, "proj_idx", cb.current()),
                                                   self.render_project()))

        row = tk.Frame(pg, bg=P["bg"])
        row.pack(fill="x")
        o1, c1 = self.card(row)
        o1.pack(side="left", fill="both", expand=True)
        self.label(c1, "Controlador", "h2").pack(anchor="w", pady=(0, 8))
        grid = tk.Frame(c1, bg=P["card"])
        grid.pack(fill="x")
        kv = [("Nombre", info.get("Name", "-")), ("Procesador", info.get("ProcessorType", "-")),
              ("Firmware", f"{info.get('MajorRev', '?')}.{info.get('MinorRev', '?')}"),
              ("Software", info.get("SoftwareRevision", "-")), ("Tipo de export", info.get("TargetType", "-")),
              ("Exportado", info.get("ExportDate", "-")), ("Última modificación", info.get("LastModifiedDate", "-")),
              ("Tareas", str(len(proj.tasks))), ("Programas", str(len(proj.programs))),
              ("Rutinas", str(sum(len(p["routines"]) for p in proj.programs.values()))),
              ("Tags", f"{proj.tag_count:,}"), ("Renglones RLL", f"{r.nrungs:,}"),
              ("Líneas ST", f"{r.nst:,}"), ("Índice de riesgo", f"{r.score:.1f}  ({r.grade})")]
        for i, (k, v) in enumerate(kv):
            rr, cc = i % 7, (i // 7) * 2
            self.label(grid, k, "small", "muted").grid(row=rr, column=cc, sticky="w", pady=3, padx=(0, 12))
            self.label(grid, v, "h3").grid(row=rr, column=cc + 1, sticky="w", pady=3, padx=(0, 30))

        o2, c2 = self.card(row, width=460)
        o2.pack(side="left", fill="both", padx=(12, 0))
        self.label(c2, "Instrucciones más usadas", "h2").pack(anchor="w", pady=(0, 8))
        top_i = r.instr.most_common(9)
        BarChart(c2, self, [(k, v, P["accent"]) for k, v in top_i], max_rows=9, row_h=22).pack(
            fill="both", expand=True)

        o3, c3 = self.card(pg)
        o3.pack(fill="both", expand=True, pady=(12, 0))
        self.label(c3, "Programas", "h2").pack(anchor="w", pady=(0, 8))
        task_of = {p: t["name"] for t in proj.tasks for p in t["programs"]}
        per_prog = Counter(f["prog"] for f in r.findings)
        per_prog_alta = Counter(f["prog"] for f in r.findings if f["sev"] == "ALTA")
        cols = ("prog", "task", "main", "routines", "rungs", "st", "find", "alta", "state")
        wrap, tv = self.tree(c3, cols, ("Programa", "Tarea", "Rutina principal", "Rutinas", "Renglones",
                                        "Líneas ST", "Hallazgos", "ALTA", "Estado"),
                             (200, 150, 150, 80, 90, 80, 90, 60, 150), height=6, stretch="prog")
        wrap.pack(fill="both", expand=True)
        for name, p in proj.programs.items():
            state = "Deshabilitado" if p["disabled"] else ("Sin tarea" if proj.tasks and name not in task_of
                                                           and not p["parent"] else "Activo")
            tags = ("ALTA",) if per_prog_alta[name] else (("MEDIA",) if state != "Activo" else ())
            tv.insert("", "end", values=(name, task_of.get(name, "-"), p["main"] or "-", len(p["routines"]),
                                         p["rungs"], p["st"], per_prog[name], per_prog_alta[name], state),
                      tags=tags)

    # -------------------------------------------------------------- reglas --
    def _build_rules_page(self, pg):
        P = self.P
        bar = tk.Frame(pg, bg=P["bg"])
        bar.pack(fill="x", pady=(0, 10))
        self.label(bar, "Doble clic o Espacio para activar / desactivar. Los cambios se aplican en el "
                   "próximo análisis.", "small", "muted", "bg").pack(side="left")
        FlatButton(bar, self, "Desactivar todas", lambda: self.set_all_rules(False), small=True).pack(
            side="right", padx=(6, 0))
        FlatButton(bar, self, "Activar todas", lambda: self.set_all_rules(True), small=True).pack(side="right")
        o, c = self.card(pg, pad=1)
        o.pack(fill="both", expand=True)
        cols = ("on", "sev", "code", "title", "cat")
        wrap, self.rtv = self.tree(c, cols, ("Activa", "Severidad", "Código", "Regla", "Categoría"),
                                   (80, 100, 230, 320, 140), height=14, stretch="title")
        wrap.pack(fill="both", expand=True)
        self.rtv.bind("<Double-1>", lambda e: self._toggle_sel_rule())
        self.rtv.bind("<space>", lambda e: self._toggle_sel_rule())
        self.rtv.bind("<<TreeviewSelect>>", lambda e: self._rule_detail())
        do, dc = self.card(pg)
        do.pack(fill="x", pady=(12, 0))
        self.r_title = self.label(dc, "Selecciona una regla", "h2")
        self.r_title.pack(anchor="w")
        self.r_desc = self.label(dc, "", "body", "muted", justify="left", wraplength=1000)
        self.r_desc.pack(anchor="w", pady=(6, 2))
        self.r_fix = self.label(dc, "", "body", "text", justify="left", wraplength=1000)
        self.r_fix.pack(anchor="w")

    def refresh_rules(self):
        tv = self.rtv
        sel = tv.selection()
        tv.delete(*tv.get_children())
        for r in sorted(RULES.values(), key=lambda r: (SEVERITY_ORDER[r.sev], r.code)):
            on = r.code in self.enabled
            tv.insert("", "end", iid=r.code, values=("✔  Sí" if on else "—  No", r.sev, r.code,
                                                     r.title, r.category),
                      tags=(r.sev,) if on else ("dim",))
        if sel and tv.exists(sel[0]):
            tv.selection_set(sel[0])

    def _rule_detail(self):
        sel = self.rtv.selection()
        if not sel:
            return
        r = RULES[sel[0]]
        self.r_title.configure(text=f"{r.title}   ·   {r.code}   ·   {r.sev}")
        self.r_desc.configure(text=r.desc)
        self.r_fix.configure(text="Qué revisar: " + r.fix)

    def _toggle_sel_rule(self):
        sel = self.rtv.selection()
        if sel:
            self.toggle_rule(sel[0], sel[0] not in self.enabled)

    def toggle_rule(self, code, on):
        (self.enabled.add if on else self.enabled.discard)(code)
        self.cfg["disabled_rules"] = sorted(set(RULES) - self.enabled)
        save_config(self.cfg)
        self.refresh_rules()
        self.refresh_badges()

    def set_all_rules(self, on):
        self.enabled = set(RULES) if on else set()
        self.cfg["disabled_rules"] = sorted(set(RULES) - self.enabled)
        save_config(self.cfg)
        self.refresh_rules()
        self.refresh_badges()

    # ------------------------------------------------------------ registro --
    def _build_log_page(self, pg):
        P = self.P
        o, c = self.card(pg, pad=1)
        o.pack(fill="both", expand=True)
        self.log_txt = tk.Text(c, bg=P["code_bg"], fg=P["text"], font=self.F["mono"], relief="flat",
                               padx=14, pady=12, highlightthickness=0, bd=0, wrap="word")
        vs = ttk.Scrollbar(c, orient="vertical", command=self.log_txt.yview)
        self.log_txt.configure(yscrollcommand=vs.set)
        vs.pack(side="right", fill="y")
        self.log_txt.pack(fill="both", expand=True)
        self.log_txt.insert("end", "\n".join(self.log_lines) + ("\n" if self.log_lines else ""))
        self.log_txt.see("end")
        self.log_txt.configure(state="disabled")

    # --------------------------------------------------------- exportacion --
    def export_menu(self):
        if not self.btn_export.enabled:
            if not self.results:
                self.toast("Primero analiza al menos un archivo", "MEDIA")
            return
        m = self.menu()
        m.add_command(label="Informe HTML (recomendado)", accelerator="Ctrl+E", command=lambda: self.export("html"))
        m.add_command(label="Excel (.xlsx)" + ("" if xlsx_available() else "  - requiere openpyxl"),
                      command=lambda: self.export("xlsx"), state="normal" if xlsx_available() else "disabled")
        m.add_command(label="CSV", command=lambda: self.export("csv"))
        m.add_command(label="JSON", command=lambda: self.export("json"))
        m.add_separator()
        m.add_checkbutton(label="Exportar solo la vista filtrada", variable=self.only_view,
                          selectcolor=self.P["accent"])
        b = self.btn_export
        m.tk_popup(b.winfo_rootx(), b.winfo_rooty() + b.winfo_height() + 2)

    def export(self, kind):
        if not self.results or self.busy:
            return
        self.cfg["export_only_view"] = bool(self.only_view.get())
        findings = [self.findings[i] for i in self._visible()] if self.only_view.get() else self.findings
        ext = {"csv": ".csv", "html": ".html", "json": ".json", "xlsx": ".xlsx"}[kind]
        names = {"csv": "CSV", "html": "Informe HTML", "json": "JSON", "xlsx": "Excel"}
        base = os.path.splitext(self.results[0].name)[0] if len(self.results) == 1 else "auditoria_l5x"
        dest = filedialog.asksaveasfilename(
            parent=self.root, title="Guardar reporte como", defaultextension=ext,
            initialdir=self.cfg.get("last_export_dir") or self.cfg.get("last_dir") or None,
            initialfile=f"{base}_{datetime.now():%Y%m%d_%H%M}{ext}", filetypes=[(names[kind], "*" + ext)])
        if not dest:
            return
        self.cfg["last_export_dir"] = os.path.dirname(dest)
        save_config(self.cfg)
        try:
            EXPORTERS[kind](dest, self.results, findings)
        except Exception as exc:
            messagebox.showerror(APP_NAME, f"No se pudo exportar:\n{exc}", parent=self.root)
            return
        self.log(f"Reporte {kind.upper()} generado ({len(findings)} hallazgos): {dest}")
        self.toast(f"Reporte guardado\n{os.path.basename(dest)}", "ok")
        if kind == "html":
            webbrowser.open("file://" + os.path.abspath(dest))

    def about(self):
        P, F = self.P, self.F
        win = tk.Toplevel(self.root, bg=P["card"])
        win.title(f"Acerca de {APP_NAME}")
        win.resizable(False, False)
        win.transient(self.root)
        body = tk.Frame(win, bg=P["card"])
        body.pack(padx=40, pady=(30, 24))
        Logo(body, self, size=64, bg=P["card"]).pack()
        self.label(body, APP_NAME, "h1").pack(pady=(14, 0))
        self.label(body, f"Versión {__version__}", "small", "muted").pack()
        self.label(body, "Auditoría estática de proyectos Studio 5000 / Logix Designer (.L5X)",
                   "body", "muted").pack(pady=(12, 16))
        tk.Frame(body, bg=P["border"], height=1).pack(fill="x")
        self.label(body, "DESARROLLADO POR", "caps", "faint").pack(pady=(16, 2))
        self.label(body, __author__, "h2").pack()
        self.label(body, f"{APP_COPYRIGHT}. Todos los derechos reservados.", "small", "muted").pack(pady=(2, 16))
        tk.Frame(body, bg=P["border"], height=1).pack(fill="x")
        runtime = f"Python {sys.version.split()[0]}  ·  Tk {tk.TkVersion}"
        runtime += "  ·  Excel " + ("disponible" if xlsx_available() else "no disponible")
        self.label(body, runtime, "small", "faint").pack(pady=(14, 18))
        FlatButton(body, self, "Cerrar", win.destroy, kind="primary").pack()
        win.bind("<Escape>", lambda e: win.destroy())
        win.update_idletasks()
        x = self.root.winfo_rootx() + (self.root.winfo_width() - win.winfo_width()) // 2
        y = self.root.winfo_rooty() + (self.root.winfo_height() - win.winfo_height()) // 3
        win.geometry(f"+{max(x, 0)}+{max(y, 0)}")
        win.grab_set()
        win.focus_set()

    def on_close(self):
        self.cfg["geometry"] = self.root.geometry()
        self.cfg["export_only_view"] = bool(self.only_view.get())
        save_config(self.cfg)
        self.root.destroy()


def run_gui(files=()):
    if tk is None:
        print("Tkinter no esta disponible en este Python. Usa el modo consola:\n"
              "    python l5x_auditor_gui.py archivo.L5X --html reporte.html")
        return 1
    if sys.platform == "win32":
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    dnd = False
    try:
        from tkinterdnd2 import TkinterDnD
        root = TkinterDnD.Tk()
        dnd = True
    except Exception:
        root = tk.Tk()
    AuditorApp(root, files, dnd=dnd)
    root.mainloop()
    return 0


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except Exception:
            pass
    return cli(argv)


if __name__ == "__main__":
    sys.exit(main())
