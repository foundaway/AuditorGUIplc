"""Pruebas del comparador de proyectos L5X (no requieren Tk)."""

import os
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import l5x_auditor_gui as m  # noqa: E402

FIX = os.path.join(os.path.dirname(__file__), "fixtures")
BASE = os.path.join(FIX, "cmp_base.L5X")
NUEVO = os.path.join(FIX, "cmp_nuevo.L5X")

# (cambio, categoria, programa / elemento, rutina, elemento): todo lo que cambia entre los dos archivos.
EXPECTED = {
    ("Modificado", "Controlador", "-", "-", "Firmware"),
    ("Modificado", "Controlador", "-", "-", "Versión de software"),
    ("Modificado", "Tareas", "-", "-", "MainTask"),
    ("Eliminado", "Programas", "Diagnostico", "-", "Diagnostico"),
    ("Agregado", "Programas", "Empaque", "-", "Empaque"),
    ("Modificado", "Rutinas", "Principal", "Mezcla", "Mezcla"),                 # FBD
    ("Agregado", "Rutinas", "Principal", "Rutina_Nueva", "Rutina_Nueva"),
    ("Eliminado", "Rutinas", "Principal", "Rutina_Vieja", "Rutina_Vieja"),
    ("Agregado", "Rungs", "Principal", "Estacion1", "Rung 1"),                  # insertado en medio
    ("Modificado", "Rungs", "Principal", "Estacion1", "Rung 4"),                # era el rung 3
    ("Eliminado", "Rungs", "Principal", "Estacion1", "Rung 4"),                 # rung 4 de la Base
    ("Modificado", "Comentarios", "Principal", "Estacion1", "Rung 2"),          # era el rung 1
    ("Modificado", "ST", "Principal", "Calculos", "Línea 2"),
    ("Modificado", "Tags", "Controlador", "-", "Paro_Emergencia"),              # descripcion
    ("Modificado", "Tags", "Controlador", "-", "Recetas"),                      # dimensiones
    ("Modificado", "Tags", "Controlador", "-", "Sensor_Alias"),                 # alias
    ("Modificado", "Tags", "Controlador", "-", "Velocidad"),                    # tipo de dato
    ("Agregado", "Tags", "Controlador", "-", "Tag_Nuevo"),
    ("Eliminado", "Tags", "Controlador", "-", "Tag_Viejo"),
    ("Modificado", "Tags", "Principal", "-", "Local_Temp"),                     # tag de programa
    ("Agregado", "AOIs", "AOI AOI_Motor", "-", "AOI_Motor"),
    ("Modificado", "AOIs", "AOI AOI_Valvula", "-", "AOI_Valvula"),              # revision
    ("Agregado", "AOIs", "AOI AOI_Valvula", "-", "Parámetro Tiempo_Falla"),
    ("Agregado", "UDTs", "-", "-", "UDT_Receta.Temperatura"),
    ("Eliminado", "UDTs", "-", "-", "UDT_Receta.Tiempo"),
    ("Eliminado", "UDTs", "-", "-", "UDT_Viejo"),
    ("Modificado", "Módulos", "-", "-", "IO_Remoto"),                           # IP y RPI
    ("Modificado", "Módulos", "-", "-", "Local"),                               # revision del CPU
    ("Agregado", "Módulos", "-", "-", "VFD_Banda"),
}


def keys(res):
    return [(c["kind"], c["cat"], c["prog"], c["rout"], c["item"]) for c in res.changes]


def find(res, cat, item, kind=None):
    hits = [c for c in res.changes if c["cat"] == cat and c["item"] == item and (kind is None or c["kind"] == kind)]
    assert len(hits) == 1, (cat, item, kind, hits)
    return hits[0]


class CompareFixtureTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.res = m.compare_files(BASE, NUEVO)

    def test_detects_exactly_the_known_changes(self):
        found = keys(self.res)
        self.assertEqual(len(found), len(set(found)), "hay cambios duplicados")
        self.assertEqual(set(found) - EXPECTED, set(), "cambios de mas (falsos positivos)")
        self.assertEqual(EXPECTED - set(found), set(), "cambios no detectados")

    def test_inserted_rung_does_not_cascade(self):
        rungs = [c for c in self.res.changes if c["cat"] == "Rungs"]
        self.assertEqual(len(rungs), 3)
        mod = find(self.res, "Rungs", "Rung 4", "Modificado")
        self.assertEqual(mod["loc"], "Rung 3 → 4")
        self.assertIn("XIC(Permiso)", mod["detail"])
        self.assertEqual(find(self.res, "Comentarios", "Rung 2")["loc"], "Rung 1 → 2")

    def test_property_details(self):
        io = find(self.res, "Módulos", "IO_Remoto")
        changed = {k: (a, b) for k, a, b, ch in io["props"] if ch}
        self.assertEqual(changed["IP"], ("192.168.1.20", "192.168.1.21"))
        self.assertEqual(changed["RPI"], ("Rack Optimization: 20 ms", "Rack Optimization: 10 ms"))
        vel = {k: (a, b) for k, a, b, ch in find(self.res, "Tags", "Velocidad")["props"] if ch}
        self.assertEqual(vel["DataType"], ("DINT", "REAL"))
        self.assertEqual(find(self.res, "Controlador", "Firmware")["detail"], "32.11 → 33.11")

    def test_values_ignored_by_default_but_available(self):
        self.assertNotIn("Contador_Piezas", [c["item"] for c in self.res.changes])
        full = m.compare_files(BASE, NUEVO, ignore_values=False)
        extra = set(keys(full)) - EXPECTED
        self.assertEqual(extra, {("Modificado", "Tags", "Controlador", "-", "Contador_Piezas")})

    def test_summary_counts(self):
        by_kind = {}
        for k, *_ in EXPECTED:
            by_kind[k] = by_kind.get(k, 0) + 1
        self.assertEqual(self.res.counts, {k: by_kind.get(k, 0) for k in m.CMP_KINDS})
        self.assertEqual(sum(sum(c.values()) for c in self.res.by_cat.values()), len(EXPECTED))

    def test_side_by_side_rows(self):
        rows = m.diff_rows(find(self.res, "Rungs", "Rung 4", "Modificado"))
        (_, left, _), (_, right, _) = rows[0]
        self.assertIn(("XIC(Permiso)", "ins"), right)
        self.assertNotIn("del", [t for _, t in left])
        st = m.diff_rows(find(self.res, "ST", "Línea 2"))
        (_, segs_l, style), (_, segs_r, _) = st[0]
        self.assertEqual(style, "chg")
        self.assertIn(("100", "del"), segs_l)
        self.assertIn(("150", "ins"), segs_r)

    def test_exports(self):
        tmp = tempfile.mkdtemp()
        html_path = os.path.join(tmp, "cmp.html")
        m.export_compare_html(html_path, self.res)
        with open(html_path, encoding="utf-8") as fh:
            body = fh.read()
        self.assertIn("Comparación de proyectos L5X", body)
        self.assertEqual(body.count("<tr data-kind="), len(EXPECTED))
        if m.xlsx_available():
            from openpyxl import load_workbook
            x = os.path.join(tmp, "cmp.xlsx")
            m.export_compare_xlsx(x, self.res)
            wb = load_workbook(x)
            self.assertEqual(wb["Cambios"].max_row - 1, len(EXPECTED))

    def test_cli(self):
        out = os.path.join(tempfile.mkdtemp(), "c.html")
        self.assertEqual(m.cli(["--compare", BASE, NUEVO, "--html", out]), 0)
        self.assertTrue(os.path.getsize(out) > 1000)


def _rewrite(src, dest, edit):
    tree = ET.parse(src)
    edit(tree.getroot())
    tree.write(dest, encoding="utf-8", xml_declaration=True)
    return dest


class RobustnessTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def test_identical_project_has_no_changes(self):
        self.assertEqual(m.compare_files(BASE, BASE).changes, [])

    def test_cosmetic_differences_are_ignored(self):
        def cosmetic(root):
            root.set("ExportDate", "Sun Jan 01 00:00:00 2034")
            for el in root.iter():
                el.attrib = dict(reversed(list(el.attrib.items())))   # otro orden de atributos
                if el.tag == "Controller":
                    el.set("LastModifiedDate", "Sun Jan 01 00:00:00 2034")
                if el.tag in ("Text", "Line") and el.text:
                    el.text = "  " + el.text.replace(")", ") ") + "  "
        other = _rewrite(BASE, os.path.join(self.tmp, "cosm.L5X"), cosmetic)
        self.assertEqual(m.compare_files(BASE, other).changes, [])

    def test_missing_sections_warn_instead_of_failing(self):
        def strip(root):
            ctrl = root.find("Controller")
            for sec in ("Modules", "DataTypes", "AddOnInstructionDefinitions"):
                ctrl.remove(ctrl.find(sec))
        old = _rewrite(BASE, os.path.join(self.tmp, "viejo.L5X"), strip)
        res = m.compare_files(old, NUEVO)
        cats = {c["cat"] for c in res.changes}
        self.assertFalse(cats & {"Módulos", "UDTs", "AOIs"})
        text = " ".join(res.warnings)
        for sec in ("Modules", "DataTypes", "AddOnInstructionDefinitions"):
            self.assertIn(sec, text)
        self.assertIn("Rungs", cats)                 # el resto se sigue comparando

    def test_partial_export(self):
        def partial(root):
            root.set("TargetType", "Program")
            ctrl = root.find("Controller")
            for sec in ("Modules", "DataTypes", "AddOnInstructionDefinitions", "Tasks", "Tags"):
                ctrl.remove(ctrl.find(sec))
        a = _rewrite(BASE, os.path.join(self.tmp, "pa.L5X"), partial)
        res = m.compare_files(a, a)
        self.assertEqual(res.changes, [])
        self.assertTrue(any("Export parcial" in w for w in res.warnings))

    def test_protected_routine(self):
        def protect(content):
            def edit(root):
                prog = root.find(".//Program[@Name='Principal']")
                routs = prog.find("Routines")
                old = routs.find("Routine[@Name='Rutina_Vieja']")
                routs.remove(old)
                enc = ET.SubElement(routs, "EncodedData", EncodedType="Routine", Name="Rutina_Vieja", Type="RLL")
                enc.text = content
            return edit
        a = _rewrite(BASE, os.path.join(self.tmp, "p1.L5X"), protect("AAAA"))
        b = _rewrite(BASE, os.path.join(self.tmp, "p2.L5X"), protect("BBBB"))
        res = m.compare_files(a, b)
        self.assertEqual(keys(res), [("Modificado", "Rutinas", "Principal", "Rutina_Vieja", "Rutina_Vieja")])
        self.assertTrue(any("protegida" in w for w in res.warnings))

    def test_large_routine_insert_is_local(self):
        rungs = [{"program": "P", "routine": "R", "number": str(i), "text": f"XIC(In_{i})OTE(Out_{i});",
                  "comment": ""} for i in range(3000)]
        new = rungs[:1500] + [{"program": "P", "routine": "R", "number": "x", "text": "XIC(Extra)OTE(Q);",
                               "comment": ""}] + rungs[1500:]
        out = []
        m._diff_rll(out, "P", "R", rungs, new)
        self.assertEqual([(c["kind"], c["detail"]) for c in out], [("Agregado", "XIC(Extra)OTE(Q);")])


if __name__ == "__main__":
    unittest.main()
